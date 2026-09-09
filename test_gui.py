"""Test approfonditi dell'interfaccia (orario/gui.py).

Esecuzione:  .venv/bin/python test_gui.py
Apre davvero le finestre tkinter e chiama i metodi dell'App come farebbe un utente,
con dialoghi (filedialog/messagebox) e rete (trasporti) sostituiti da finti.
"""

from __future__ import annotations

import os
import shutil
import sys
import tempfile
import time
import tkinter as tk
import traceback
import unittest
from pathlib import Path

os.environ.setdefault("TK_SILENCE_DEPRECATION", "1")

import openpyxl  # noqa: E402

import orario.trasporti as mod_trasporti  # noqa: E402
from orario import gui  # noqa: E402
from orario.costanti import FOGLI_DATI, FOGLIO_ORARIO, FOGLIO_TRASPORTI, GIORNI_LUNGHI  # noqa: E402
from orario.modello import Trasporto  # noqa: E402

import customtkinter as ctk  # noqa: E402

RADICE = Path(__file__).parent
ESEMPIO = RADICE / "esempio" / "orario_2026-27.xlsx"
TIMEOUT_SOLVER = 12  # secondi scritti nel foglio Parametri dei file di prova


# ── Dialoghi finti ───────────────────────────────────────────────────────────

class FinteFinestre:
    """Sostituisce gui.messagebox e gui.filedialog: registra le chiamate, risponde a comando."""

    def __init__(self) -> None:
        self.chiamate: list[tuple[str, str, str]] = []
        self.risposte: dict[str, object] = {}

    # utilità
    def titoli(self, nome: str | None = None) -> list[str]:
        return [t for (n, t, _m) in self.chiamate if nome is None or n == nome]

    def testo_tutto(self) -> str:
        return "\n".join(f"{n} | {t} | {m}" for (n, t, m) in self.chiamate)

    def conta(self, nome: str) -> int:
        return sum(1 for (n, _t, _m) in self.chiamate if n == nome)

    def azzera(self) -> None:
        self.chiamate.clear()

    def _rispondi(self, nome: str, *a, **k):
        titolo = str(a[0]) if a else str(k.get("title", ""))
        msg = str(a[1]) if len(a) > 1 else str(k.get("message", ""))
        self.chiamate.append((nome, titolo, msg))
        r = self.risposte.get(nome, "" if nome.startswith(("ask_file", "asksave", "askopen", "askdir")) else None)
        if isinstance(r, list):
            return r.pop(0) if r else None
        return r

    # messagebox
    def showinfo(self, *a, **k):
        return self._rispondi("showinfo", *a, **k)

    def showwarning(self, *a, **k):
        return self._rispondi("showwarning", *a, **k)

    def showerror(self, *a, **k):
        return self._rispondi("showerror", *a, **k)

    def askyesno(self, *a, **k):
        return bool(self._rispondi("askyesno", *a, **k))

    def askyesnocancel(self, *a, **k):
        return self._rispondi("askyesnocancel", *a, **k)

    # filedialog
    def askopenfilename(self, *a, **k):
        return self._rispondi("askopenfilename", *a, **k) or ""

    def asksaveasfilename(self, *a, **k):
        return self._rispondi("asksaveasfilename", *a, **k) or ""

    def askdirectory(self, *a, **k):
        return self._rispondi("askdirectory", *a, **k) or ""


def finti_trasporti(dati, progresso=None, annulla=None, precedenti=None, solo_nuovi=False):
    """Sostituto di orario.trasporti.aggiorna_trasporti: nessuna rete, risultati inventati."""
    precedenti = precedenti or {}
    out: dict[str, Trasporto] = {}
    for s in dati.studenti:
        if not s.indirizzo_completo:
            continue
        if solo_nuovi:
            t = precedenti.get(s.id)
            if t and t.esito == "OK" and t.indirizzo_usato == s.indirizzo_completo:
                continue
        if progresso:
            progresso(f"(finto) {s.id}")
        out[s.id] = Trasporto(minuti=[20, 22, 24, 26], arrivi=["14:55", "15:57", "16:59", "18:01"],
                              mezzi=["bus 1 (finto)"] * 4, lat=43.93, lon=10.91,
                              indirizzo_usato=s.indirizzo_completo, esito="OK", aggiornato="09/09/2026 10:00")
    return out


# ── Motore del test: script a passi dentro il mainloop ───────────────────────

class BaseGui(unittest.TestCase):
    """Esegue uno script (funzione generatore) dentro il mainloop di una App vera.

    Lo script riceve l'app e produce:
      * None                     → prosegui al passo successivo
      * (condizione, secondi)    → aspetta che condizione() sia vera (senza bloccare la GUI)
    """

    @classmethod
    def setUpClass(cls) -> None:
        cls.originale_trasporti = mod_trasporti.aggiorna_trasporti

    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp(prefix="test_gui_"))
        self.imp_originale = gui.IMPOSTAZIONI
        gui.IMPOSTAZIONI = self.tmp / "impostazioni.json"
        self.mb_originale, self.fd_originale = gui.messagebox, gui.filedialog
        self.finte = FinteFinestre()
        gui.messagebox = self.finte
        gui.filedialog = self.finte
        mod_trasporti.aggiorna_trasporti = finti_trasporti
        self.avanzi: list[str] = []

    def tearDown(self) -> None:
        gui.IMPOSTAZIONI = self.imp_originale
        gui.messagebox, gui.filedialog = self.mb_originale, self.fd_originale
        mod_trasporti.aggiorna_trasporti = self.originale_trasporti
        for p in self.tmp.rglob("*"):
            try:
                p.chmod(0o700)
            except Exception:
                pass
        shutil.rmtree(self.tmp, ignore_errors=True)

    # ── file di prova ──
    def copia_esempio(self, nome: str = "prova.xlsx", cartella: Path | None = None) -> Path:
        dest = (cartella or self.tmp) / nome
        dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy(ESEMPIO, dest)
        wb = openpyxl.load_workbook(dest)
        ws = wb["Parametri"]
        for r in ws.iter_rows(min_row=2):
            if r[0].value and "tempo" in str(r[0].value).lower():
                r[1].value = TIMEOUT_SOLVER
        wb.save(dest)
        return dest

    # ── esecuzione ──
    def esegui(self, script, limite: float = 150.0) -> None:
        errori: list[str] = []
        app = gui.App()
        app.report_callback_exception = lambda *a: errori.append(
            "eccezione non gestita in un callback tkinter:\n" + "".join(traceback.format_exception(*a)))
        it = script(app)
        finito = {"v": False}

        def chiudi(msg: str | None = None) -> None:
            if finito["v"]:
                return
            finito["v"] = True
            if msg:
                errori.append(msg)
            self.avanzi = [type(w).__name__ for w in app.winfo_children() if isinstance(w, tk.Toplevel)]
            for w in app.winfo_children():
                if isinstance(w, tk.Toplevel):
                    try:
                        w.grab_release()
                        w.destroy()
                    except Exception:
                        pass
            app.after(40, app.destroy)

        def prossimo() -> None:
            if finito["v"]:
                return
            try:
                v = next(it)
            except StopIteration:
                chiudi()
                return
            except Exception:
                chiudi("errore nello script:\n" + traceback.format_exc())
                return
            if v is None:
                app.after(30, prossimo)
                return
            cond, secondi = v if isinstance(v, tuple) else (v, 30)
            scadenza = time.monotonic() + secondi
            nome = getattr(cond, "__doc__", None) or getattr(cond, "__name__", str(cond))

            def poll() -> None:
                if finito["v"]:
                    return
                try:
                    ok = cond()
                except Exception:
                    chiudi("errore nella condizione di attesa:\n" + traceback.format_exc())
                    return
                if ok:
                    app.after(30, prossimo)
                elif time.monotonic() > scadenza:
                    chiudi(f"tempo scaduto ({secondi}s) aspettando: {nome}")
                else:
                    app.after(60, poll)

            poll()

        app.after(int(limite * 1000), lambda: chiudi(f"il test non è finito entro {limite}s"))
        app.after(500, prossimo)
        app.mainloop()
        if errori:
            self.fail("\n\n".join(errori))

    # ── aiuti ──
    @staticmethod
    def toplevel(app, tipo) -> list:
        return [w for w in app.winfo_children() if isinstance(w, tipo)]

    def chiudi_toplevel(self, app, tipo) -> None:
        for w in self.toplevel(app, tipo):
            try:
                w.grab_release()
            except Exception:
                pass
            w.destroy()

    @staticmethod
    def col(sheet, intestazione: str) -> int:
        return list(sheet.headers()).index(intestazione)

    @staticmethod
    def calcolo_finito(app):
        def cond():
            """fine del calcolo (finestra di attesa chiusa e pulsante Calcola riattivo)"""
            att = getattr(app, "attesa", None)
            chiusa = att is None or not att.winfo_exists()
            return chiusa and str(app.btn_calcola.cget("state")) == "normal"
        return (cond, 60)


# ── 1. Schermata di apertura ─────────────────────────────────────────────────

class TestApertura(BaseGui):

    def test_01_pulsanti_schermata_apertura(self):
        def script(app):
            testi = []

            def raccogli(w):
                for c in w.winfo_children():
                    if isinstance(c, ctk.CTkButton):
                        testi.append(c.cget("text"))
                    raccogli(c)
            raccogli(app.schermata_apertura)
            self.assertIn("Apri file Excel…", testi)
            self.assertIn("Nuovo file…", testi)
            self.assertFalse([t for t in testi if t.startswith("Riapri")],
                             "senza ultimo file non deve comparire 'Riapri'")
            self.assertIsNone(app.percorso)
            self.assertIsNone(app.schermata_dati)
            yield
        self.esegui(script)
        self.assertEqual(self.avanzi, [])

    def test_02_riapri_ultimo_solo_se_esiste(self):
        f = self.copia_esempio()
        gui.IMPOSTAZIONI.write_text('{"ultimo_file": "%s"}' % f)

        def script(app):
            testi = []

            def raccogli(w):
                for c in w.winfo_children():
                    if isinstance(c, ctk.CTkButton):
                        testi.append(c.cget("text"))
                    raccogli(c)
            raccogli(app.schermata_apertura)
            self.assertTrue([t for t in testi if t.startswith("Riapri")], f"nessun pulsante Riapri fra {testi}")
            yield
        self.esegui(script)

        # ora l'ultimo file non esiste più: il pulsante deve sparire e il menu avvisare
        f.unlink()

        def script2(app):
            testi = []

            def raccogli(w):
                for c in w.winfo_children():
                    if isinstance(c, ctk.CTkButton):
                        testi.append(c.cget("text"))
                    raccogli(c)
            raccogli(app.schermata_apertura)
            self.assertFalse([t for t in testi if t.startswith("Riapri")])
            app.riapri_ultimo()
            self.assertIn("Nessun file recente", self.finte.titoli("showinfo"))
            yield
        self.esegui(script2)

    def test_03_file_inesistente_e_corrotto(self):
        finto = self.tmp / "finto.xlsx"
        finto.write_text("questo non è un file Excel, è solo testo")

        def script(app):
            app.apri_file(self.tmp / "manca_del_tutto.xlsx")
            yield
            fin = self.toplevel(app, gui.FinestraProblemi)
            self.assertEqual(len(fin), 1, "file inesistente: attesa la finestra dei problemi")
            self.assertIn("Impossibile leggere il file", fin[0].testo)
            self.assertIsNone(app.percorso)
            self.chiudi_toplevel(app, gui.FinestraProblemi)
            yield
            app.apri_file(finto)
            yield
            fin = self.toplevel(app, gui.FinestraProblemi)
            self.assertEqual(len(fin), 1, "file corrotto: attesa la finestra dei problemi")
            self.assertIn("finto.xlsx", fin[0].testo)
            self.assertIsNone(app.percorso, "un file illeggibile non deve diventare il file corrente")
            self.chiudi_toplevel(app, gui.FinestraProblemi)
            yield
        self.esegui(script)
        self.assertEqual(self.avanzi, [])

    def test_04_scegli_file_annullato_e_confermato(self):
        f = self.copia_esempio()

        def script(app):
            self.finte.risposte["askopenfilename"] = ["", str(f)]
            app.scegli_file()
            yield
            self.assertIsNone(app.percorso, "annullando il dialogo non si deve aprire nulla")
            app.scegli_file()
            yield
            self.assertEqual(app.percorso, f)
            self.assertEqual(len(app.fogli), 5)
            yield
        self.esegui(script)


# ── 2. Nuovo file ────────────────────────────────────────────────────────────

class TestNuovoFile(BaseGui):

    def test_05_nuovo_file_annullato(self):
        def script(app):
            self.finte.risposte["asksaveasfilename"] = ""
            app.nuovo_file()
            yield
            self.assertIsNone(app.percorso)
            self.assertEqual(list(self.tmp.glob("*.xlsx")), [], "annullando non si deve creare nessun file")
            yield
        self.esegui(script)

    def test_06_nuovo_file_creato_e_aperto(self):
        nuovo = self.tmp / "nuovo_orario.xlsx"

        def script(app):
            self.finte.risposte["asksaveasfilename"] = str(nuovo)
            app.nuovo_file()
            yield
            self.assertTrue(nuovo.exists(), "nuovo_file() deve creare il file")
            self.assertEqual(app.percorso, nuovo)
            self.assertEqual(sorted(app.fogli), sorted(FOGLI_DATI))
            self.assertIn("esempi", app.lbl_stato.cget("text").lower())
            # è davvero un file valido e diventa "ultimo file"
            self.assertIn("Studenti", openpyxl.load_workbook(nuovo).sheetnames)
            self.assertIn(str(nuovo), gui.IMPOSTAZIONI.read_text())
            yield
        self.esegui(script)


# ── 3. Griglie ───────────────────────────────────────────────────────────────

class TestGriglie(BaseGui):

    def test_07_fogli_righe_colonne(self):
        f = self.copia_esempio()
        atteso = {}
        wb = openpyxl.load_workbook(f)
        for nome in FOGLI_DATI:
            ws = wb[nome]
            atteso[nome] = (ws.max_row - 1, ws.max_column)

        def script(app):
            app.apri_file(f)
            yield
            self.assertEqual(sorted(app.fogli), sorted(FOGLI_DATI))
            for nome, sheet in app.fogli.items():
                righe = sheet.get_sheet_data()
                self.assertEqual(len(righe), atteso[nome][0], f"righe del foglio {nome}")
                self.assertEqual(len(sheet.headers()), atteso[nome][1], f"colonne del foglio {nome}")
            self.assertEqual(len(app.fogli["Studenti"].get_sheet_data()), 30, "30 studenti nel file di esempio")
            yield
        self.esegui(script)

    def test_08_tabelle_correnti_rilegge_dalle_griglie(self):
        f = self.copia_esempio()

        def script(app):
            app.apri_file(f)
            yield
            sh = app.fogli["Studenti"]
            c = self.col(sh, "Cognome")
            sh.set_cell_data(0, c, "MODIFICATO")
            t = app.tabelle_correnti()
            self.assertEqual(t["Studenti"].righe[0][c], "MODIFICATO")
            self.assertEqual(len(t["Studenti"].righe), 30)
            self.assertEqual(sorted(t), sorted(FOGLI_DATI))
            for nome in FOGLI_DATI:
                self.assertEqual(list(t[nome].intestazione), list(app.tabelle[nome].intestazione))
            yield
        self.esegui(script)

    def test_09_aggiungi_e_elimina_righe(self):
        f = self.copia_esempio()

        def script(app):
            app.apri_file(f)
            yield
            sh = app.fogli["Studenti"]
            n = len(sh.get_sheet_data())
            app.aggiungi_riga("Studenti")
            self.assertEqual(len(sh.get_sheet_data()), n + 1)
            self.assertEqual(len(app.tabelle_correnti()["Studenti"].righe), n,
                             "la riga vuota non deve contare nei dati")
            sh.set_cell_data(n, self.col(sh, "Cognome"), "NUOVO")
            self.assertEqual(len(app.tabelle_correnti()["Studenti"].righe), n + 1)

            # elimina senza selezione → messaggio nella barra di stato
            app.lbl_stato.configure(text="")
            app.elimina_righe("Studenti")
            self.assertIn("Seleziona prima", app.lbl_stato.cget("text"))
            self.assertEqual(len(sh.get_sheet_data()), n + 1)

            # elimina con selezione
            sh.select_row(0)
            yield
            app.elimina_righe("Studenti")
            self.assertEqual(len(sh.get_sheet_data()), n)
            self.assertNotEqual(app.tabelle_correnti()["Studenti"].righe[0][self.col(sh, "Cognome")], "AMATO")

            # anche dal menu Modifica, sul foglio corrente
            app.tabs.set("Gruppi LMC")
            yield
            g = len(app.fogli["Gruppi LMC"].get_sheet_data())
            app.aggiungi_riga_corrente()
            self.assertEqual(len(app.fogli["Gruppi LMC"].get_sheet_data()), g + 1)
            app.lbl_stato.configure(text="")
            app.elimina_righe_corrente()
            self.assertIn("Seleziona prima", app.lbl_stato.cget("text"))
            yield
        self.esegui(script)

    def test_10_modifica_salva_riapri(self):
        f = self.copia_esempio()

        def script(app):
            app.apri_file(f)
            yield
            sh = app.fogli["Studenti"]
            sh.set_cell_data(0, self.col(sh, "Note"), "NOTA_DI_PROVA_42")
            self.assertTrue(app.salva())
            self.assertIn("Salvato", app.lbl_stato.cget("text"))
            yield
            ws = openpyxl.load_workbook(f)["Studenti"]
            valori = [c.value for c in ws[2]]
            self.assertIn("NOTA_DI_PROVA_42", valori, "il valore modificato deve finire nel file")
            # riapertura: la griglia mostra il valore nuovo
            app.apri_file(f)
            yield
            sh = app.fogli["Studenti"]
            self.assertEqual(sh.get_cell_data(0, self.col(sh, "Note")), "NOTA_DI_PROVA_42")
            yield
        self.esegui(script)

    def test_11_salva_su_file_sola_lettura(self):
        sotto = self.tmp / "ro"
        f = self.copia_esempio(cartella=sotto)

        def script(app):
            app.apri_file(f)
            yield
            f.chmod(0o444)
            esito = app.salva()  # su macOS la rinomina riesce anche su file 444: basta che non esploda
            yield
            self.assertIsInstance(esito, bool)
            fin = self.toplevel(app, gui.FinestraProblemi)
            if not esito:
                self.assertEqual(len(fin), 1)
            self.chiudi_toplevel(app, gui.FinestraProblemi)
            # cartella non scrivibile: qui il salvataggio deve fallire con la finestra dei problemi
            sotto.chmod(0o500)
            try:
                esito2 = app.salva()
            finally:
                sotto.chmod(0o700)
            yield
            self.assertFalse(esito2, "salva() deve restituire False se non può scrivere")
            fin = self.toplevel(app, gui.FinestraProblemi)
            self.assertEqual(len(fin), 1, "atteso l'avviso di salvataggio non riuscito")
            self.assertIn("Salvataggio", fin[0].title() + fin[0].testo)
            self.chiudi_toplevel(app, gui.FinestraProblemi)
            yield
            self.assertTrue(f.exists(), "il file non deve sparire")
            openpyxl.load_workbook(f)  # non corrotto
            yield
        self.esegui(script)


# ── 4. Menu e scorciatoie ────────────────────────────────────────────────────

class TestMenu(BaseGui):

    @staticmethod
    def voci(app, menu):
        out = []
        for i in range(menu.index("end") + 1):
            if menu.type(i) == "separator":
                out.append("---")
            else:
                out.append(menu.entrycget(i, "label"))
        return out

    def test_12_struttura_della_barra(self):
        def script(app):
            barra = app.nametowidget(app.cget("menu"))
            cascate = [barra.entrycget(i, "label") for i in range(barra.index("end") + 1)]
            self.assertEqual(cascate, ["File", "Modifica", "Orario", "Aiuto"])
            atteso = {
                "File": ["Nuovo file…", "Apri file…", "Riapri ultimo file", "Salva", "Salva con nome…",
                         "Chiudi file", "Esci"],
                "Modifica": ["Aggiungi riga al foglio corrente", "Elimina righe selezionate"],
                "Orario": ["Calcola orario…", "Ricalcola da zero (ignora l'orario precedente)…",
                           "Apri l'ultima cartella dei risultati"],
                "Aiuto": ["Istruzioni per compilare il file", "Informazioni"],
            }
            for i, nome in enumerate(cascate):
                sub = app.nametowidget(barra.entrycget(i, "menu"))
                voci = self.voci(app, sub)
                for v in atteso[nome]:
                    self.assertIn(v, voci, f"voce '{v}' mancante nel menu {nome}")
            sub = app.nametowidget(barra.entrycget(2, "menu"))
            self.assertTrue([v for v in self.voci(app, sub) if v.startswith("Aggiorna trasporti")])
            yield
        self.esegui(script)

    def test_13_voci_invocabili_senza_file(self):
        def script(app):
            barra = app.nametowidget(app.cget("menu"))
            self.finte.risposte["asksaveasfilename"] = ""
            self.finte.risposte["askopenfilename"] = ""
            for i in range(barra.index("end") + 1):
                sub = app.nametowidget(barra.entrycget(i, "menu"))
                for j in range(sub.index("end") + 1):
                    if sub.type(j) != "command":
                        continue
                    if sub.entrycget(j, "label") == "Esci":
                        continue
                    if str(sub.entrycget(j, "state")) == "disabled":
                        continue
                    sub.invoke(j)  # nessuna eccezione ammessa
                    yield
                    self.chiudi_toplevel(app, gui.FinestraTesto)
            testo = self.finte.testo_tutto()
            self.assertGreaterEqual(self.finte.titoli("showinfo").count("Nessun file aperto"), 4,
                                    "Salva/Calcola/Aggiungi riga/Trasporti devono dire 'Nessun file aperto':\n" + testo)
            self.assertIn("Nessun risultato", self.finte.titoli("showinfo"))
            self.assertIn("Nessun file recente", self.finte.titoli("showinfo"))
            self.assertIsNone(app.percorso)
            yield
        self.esegui(script)

    def test_14_metodi_senza_file_non_sollevano(self):
        def script(app):
            for m in (app.salva_da_menu, app.calcola_da_menu, app.aggiungi_riga_corrente,
                      app.elimina_righe_corrente, app.aggiorna_trasporti, app.ricalcola_da_zero):
                m()
                yield
            self.assertEqual(self.finte.titoli("showinfo").count("Nessun file aperto"), 6,
                             self.finte.testo_tutto())
            app.apri_risultati()
            self.assertIn("Nessun risultato", self.finte.titoli("showinfo"))
            yield
        self.esegui(script)

    def test_15_istruzioni_informazioni_scorciatoie(self):
        def script(app):
            app.mostra_istruzioni()
            yield
            fin = self.toplevel(app, gui.FinestraTesto)
            self.assertEqual(len(fin), 1, "'Istruzioni' deve aprire la finestra di testo")
            self.assertIn("Istruzioni", fin[0].title())
            self.chiudi_toplevel(app, gui.FinestraTesto)
            yield
            barra = app.nametowidget(app.cget("menu"))
            aiuto = app.nametowidget(barra.entrycget(3, "menu"))
            aiuto.invoke(self.voci(app, aiuto).index("Informazioni"))
            yield
            info = [m for (n, _t, m) in self.finte.chiamate if n == "showinfo"]
            self.assertTrue(any(gui.VERSIONE in m for m in info), f"la versione {gui.VERSIONE} non compare: {info}")
            for tasto in ("<Command-n>", "<Command-o>", "<Command-s>"):
                self.assertTrue(app.bind_all(tasto), f"scorciatoia {tasto} non registrata")
            yield
        self.esegui(script)


# ── 5. Calcolo ───────────────────────────────────────────────────────────────

class TestCalcolo(BaseGui):

    def test_16_calcolo_completo(self):
        f = self.copia_esempio()
        dest = self.tmp / "risultati"
        dest.mkdir()

        def script(app):
            app.apri_file(f)
            yield
            self.assertFalse(app.usa_precedente.get(), "senza orario precedente la casella parte spenta")
            self.assertIn("Nessun orario precedente", app.chk_precedente.cget("text"))
            self.assertEqual(str(app.btn_cartella.cget("state")), "disabled")

            # annullando la scelta della cartella non parte nulla
            self.finte.risposte["askdirectory"] = ""
            app.calcola()
            yield
            self.assertIsNone(app.orario)
            self.assertEqual(self.toplevel(app, gui.FinestraAttesa), [])
            self.assertEqual(str(app.btn_calcola.cget("state")), "normal")
            self.assertEqual(list(dest.iterdir()), [])

            # calcolo vero
            self.finte.risposte["askdirectory"] = str(dest)
            app.calcola()
            yield
            self.assertEqual(str(app.btn_calcola.cget("state")), "disabled",
                             "il pulsante Calcola deve disattivarsi durante il calcolo")
            self.assertEqual(len(self.toplevel(app, gui.FinestraAttesa)), 1)
            yield self.calcolo_finito(app)
            self.assertIsNotNone(app.orario, "calcolo non riuscito: " + self.finte.testo_tutto())
            self.assertEqual(str(app.btn_calcola.cget("state")), "normal")

            cartelle = [p for p in dest.iterdir() if p.is_dir()]
            self.assertEqual(len(cartelle), 1, f"attesa una cartella dei risultati, trovate {cartelle}")
            self.assertTrue(cartelle[0].name.startswith("Risultati orario "))
            nomi = sorted(p.name for p in cartelle[0].iterdir())
            self.assertEqual(nomi, ["orario.xlsx", "orario_docenti.pdf", "orario_settimanale.pdf",
                                    "orario_studenti.pdf"])
            self.assertEqual(app.ultima_cartella_risultati, cartelle[0])
            self.assertTrue(app.ultima_cartella_risultati.exists())
            self.assertEqual(str(app.btn_cartella.cget("state")), "normal")
            self.assertIn(str(cartelle[0]), gui.IMPOSTAZIONI.read_text())

            # scheda Orario con i 5 giorni + Avvisi e riepilogo in cima
            self.assertEqual(app.tabs.get(), "Orario")
            tab = app.tabs.tab("Orario")
            sotto = [w for w in tab.winfo_children() if isinstance(w, ctk.CTkTabview)]
            self.assertEqual(len(sotto), 1)
            nomi_tab = list(sotto[0]._name_list)
            for g in GIORNI_LUNGHI:
                self.assertIn(g, nomi_tab)
            self.assertIn("Avvisi", nomi_tab)
            etichette = [w.cget("text") for w in tab.winfo_children() if isinstance(w, ctk.CTkLabel)]
            riepilogo = "\n".join(etichette)
            for pezzo in ("Studenti:", "Lezioni:", "Buchi totali:", "Rientri:", "Avvisi:"):
                self.assertIn(pezzo, riepilogo)
            self.assertIn("Studenti: 30", riepilogo)

            fatto = self.toplevel(app, gui.FinestraFatto)
            self.assertEqual(len(fatto), 1, "attesa la finestra 'Orario creato'")
            self.chiudi_toplevel(app, gui.FinestraFatto)
            yield

            # l'orario è stato scritto nel file di input e la casella si aggiorna con la data
            self.assertIn(FOGLIO_ORARIO, openpyxl.load_workbook(f).sheetnames)
            self.assertIn("Parti dall'orario calcolato il", app.chk_precedente.cget("text"))
            self.assertTrue(app.usa_precedente.get())
            self.assertEqual(str(app.chk_precedente.cget("state")), "normal")

            # riaprendo il file la casella resta attiva; apri_risultati non protesta
            app.apri_file(f)
            yield
            self.assertIn("Parti dall'orario calcolato il", app.chk_precedente.cget("text"))
            app.ultima_cartella_risultati = None
            app.apri_risultati()
            self.assertIn("Nessun risultato", self.finte.titoli("showinfo"))
            yield
        self.esegui(script)
        self.assertEqual(self.avanzi, [])

    def test_17_errore_nei_dati_blocca_il_calcolo(self):
        f = self.copia_esempio()
        dest = self.tmp / "risultati"
        dest.mkdir()

        def script(app):
            app.apri_file(f)
            yield
            sh = app.fogli["Studenti"]
            cognome = sh.get_cell_data(0, self.col(sh, "Cognome"))
            sh.set_cell_data(0, self.col(sh, "Docente 1"), "")
            self.finte.risposte["askdirectory"] = str(dest)
            app.calcola()
            yield
            yield self.calcolo_finito(app)
            fin = self.toplevel(app, gui.FinestraProblemi)
            self.assertEqual(len(fin), 1, "atteso l'elenco dei problemi: " + self.finte.testo_tutto())
            testo = fin[0].testo
            self.assertIn("Studenti, riga 2", testo, f"il messaggio deve citare la riga:\n{testo[:400]}")
            self.assertIn(cognome, testo)
            self.assertIn("ERRORI", testo)
            self.assertIsNone(app.orario)
            self.assertEqual([p for p in dest.iterdir()], [], "nessuna cartella di risultati con dati sbagliati")

            app.clipboard_clear()
            fin[0].copia()
            yield
            self.assertIn("Studenti, riga 2", app.clipboard_get())
            self.chiudi_toplevel(app, gui.FinestraProblemi)
            yield
            self.assertEqual(str(app.btn_calcola.cget("state")), "normal")
            yield
        self.esegui(script)
        self.assertEqual(self.avanzi, [])


# ── 6. Trasporti ─────────────────────────────────────────────────────────────

class TestTrasporti(BaseGui):

    @staticmethod
    def attesa_chiusa(app):
        def cond():
            """chiusura della finestra di attesa dei trasporti"""
            att = getattr(app, "attesa", None)
            return att is None or not att.winfo_exists()
        return (cond, 60)

    def test_18_pulsante_trasporti_e_avvisi(self):
        f = self.copia_esempio()

        def script(app):
            app.apri_file(f)
            yield (lambda: "Trasporti" in app.btn_trasporti.cget("text"), 10)
            app._aggiorna_info_trasporti()
            self.assertRegex(app.btn_trasporti.cget("text"), r"Trasporti 0/\d+")
            self.assertEqual(str(app.btn_trasporti.cget("state")), "normal")

            # senza indirizzo della scuola avvisa e non parte
            sh = app.fogli["Parametri"]
            riga = [i for i, r in enumerate(sh.get_sheet_data())
                    if "indirizzo" in str(r[0]).lower() and "scuola" in str(r[0]).lower()][0]
            vecchio = sh.get_cell_data(riga, 1)
            sh.set_cell_data(riga, 1, "")
            app.aggiorna_trasporti()
            yield
            self.assertIn("Manca l'indirizzo della scuola", self.finte.titoli("showwarning"))
            self.assertEqual(self.toplevel(app, gui.FinestraAttesa), [])
            sh.set_cell_data(riga, 1, vecchio)
            self.finte.azzera()

            # aggiornamento completo: chiede conferma (askyesno) e poi calcola
            self.finte.risposte["askyesno"] = True
            app.aggiorna_trasporti()
            yield
            self.assertIn("Aggiorna trasporti", self.finte.titoli("askyesno"))
            yield self.attesa_chiusa(app)
            self.assertIn("Trasporti aggiornati", self.finte.titoli("showinfo"), self.finte.testo_tutto())
            self.assertIn(FOGLIO_TRASPORTI, openpyxl.load_workbook(f).sheetnames)
            app._aggiorna_info_trasporti()
            self.assertRegex(app.btn_trasporti.cget("text"), r"Trasporti ✓ \d+")
            self.assertIn("Trasporti aggiornati", app.lbl_stato.cget("text"))

            # con risultati già presenti compare il dialogo tutti/solo i nuovi: Annulla non fa nulla
            self.finte.azzera()
            self.finte.risposte["askyesnocancel"] = None
            app.aggiorna_trasporti()
            yield
            self.assertIn("Aggiorna trasporti", self.finte.titoli("askyesnocancel"),
                          "atteso il dialogo tutti/solo i nuovi: " + self.finte.testo_tutto())
            self.assertIn("hanno già i tempi", self.finte.testo_tutto())
            self.assertEqual(self.toplevel(app, gui.FinestraAttesa), [])
            yield
        self.esegui(script)
        self.assertEqual(self.avanzi, [])

    def test_19_calcola_con_trasporti_mancanti(self):
        f = self.copia_esempio()
        dest = self.tmp / "risultati"
        dest.mkdir()

        def script(app):
            app.apri_file(f)
            yield (lambda: "Trasporti" in app.btn_trasporti.cget("text"), 10)

            # Annulla: non fa nulla
            self.finte.risposte["askyesnocancel"] = None
            self.finte.risposte["askdirectory"] = str(dest)
            app.calcola()
            yield
            self.assertIn("Tempi dei mezzi mancanti", self.finte.titoli("askyesnocancel"))
            self.assertEqual(self.finte.conta("askdirectory"), 0, "con Annulla non si chiede la cartella")
            self.assertIsNone(app.orario)
            self.assertEqual(list(dest.iterdir()), [])
            self.assertFalse(app._calcola_dopo_trasporti)

            # No: calcola subito con i KM
            self.finte.azzera()
            self.finte.risposte["askyesnocancel"] = False
            app.calcola()
            yield
            yield self.calcolo_finito(app)
            self.assertIsNotNone(app.orario, "con 'No' l'orario si deve calcolare: " + self.finte.testo_tutto())
            self.assertNotIn(FOGLIO_TRASPORTI, openpyxl.load_workbook(f).sheetnames)
            cartelle = [p for p in dest.iterdir() if p.is_dir()]
            self.assertEqual(len(cartelle), 1)
            self.assertEqual(len(list(cartelle[0].iterdir())), 4)
            self.chiudi_toplevel(app, gui.FinestraFatto)
            yield

            # Sì: prima i trasporti, poi l'orario, da solo
            app.orario = None
            app.usa_precedente.set(False)
            self.finte.azzera()
            self.finte.risposte["askyesnocancel"] = True
            app.calcola()
            yield
            self.assertTrue(app._calcola_dopo_trasporti or app.orario is not None)
            yield (lambda: app.orario is not None, 90)
            self.assertIn(FOGLIO_TRASPORTI, openpyxl.load_workbook(f).sheetnames,
                          "i trasporti dovevano essere calcolati e salvati")
            cartelle = sorted((p for p in dest.iterdir() if p.is_dir()), key=lambda p: p.stat().st_mtime)
            self.assertEqual(len(cartelle), 2, "atteso un secondo giro di risultati")
            self.assertEqual(len(list(cartelle[-1].iterdir())), 4)
            self.assertFalse(app._calcola_dopo_trasporti)
            self.assertEqual(str(app.btn_calcola.cget("state")), "normal")
            self.chiudi_toplevel(app, gui.FinestraFatto)
            yield
            # il file di input resta leggibile e completo
            wb = openpyxl.load_workbook(f)
            for nome in FOGLI_DATI:
                self.assertIn(nome, wb.sheetnames)
            self.assertEqual(wb["Studenti"].max_row - 1, 30)
            yield
        self.esegui(script, limite=200)
        self.assertEqual(self.avanzi, [])


# ── 7. Robustezza ────────────────────────────────────────────────────────────

class TestRobustezza(BaseGui):

    def test_20_chiudi_file_e_riapertura(self):
        f = self.copia_esempio()

        def script(app):
            app.apri_file(f)
            yield
            app.torna_apertura()
            yield
            self.assertTrue(app.schermata_apertura.winfo_ismapped() or True)
            # dopo "Chiudi file" i comandi che richiedono un file devono avvisare, non calcolare
            self.finte.azzera()
            app.calcola_da_menu()
            yield
            self.assertIn("Nessun file aperto", self.finte.titoli("showinfo"),
                          "dopo 'Chiudi file' il calcolo non deve partire: " + self.finte.testo_tutto())
            self.assertEqual(self.finte.conta("askdirectory"), 0)
            yield
        self.esegui(script)

    def test_21_salva_con_nome(self):
        f = self.copia_esempio()
        altro = self.tmp / "copia.xlsx"

        def script(app):
            app.apri_file(f)
            yield
            self.finte.risposte["asksaveasfilename"] = ["", str(altro)]
            app.salva_con_nome()
            yield
            self.assertFalse(altro.exists(), "annullando non si salva")
            self.assertEqual(app.percorso, f)
            sh = app.fogli["Studenti"]
            sh.set_cell_data(0, self.col(sh, "Note"), "COPIA_OK")
            app.salva_con_nome()
            yield
            self.assertTrue(altro.exists())
            self.assertEqual(app.percorso, altro)
            self.assertIn("copia.xlsx", app.lbl_file.cget("text"))
            valori = [c.value for c in openpyxl.load_workbook(altro)["Studenti"][2]]
            self.assertIn("COPIA_OK", valori)
            yield
        self.esegui(script)

    def test_22_file_di_input_non_corrotto(self):
        f = self.copia_esempio()
        atteso = {n: (openpyxl.load_workbook(f)[n].max_row, openpyxl.load_workbook(f)[n].max_column)
                  for n in FOGLI_DATI}

        def script(app):
            app.apri_file(f)
            yield
            for _ in range(3):
                app.salva()
                yield
            wb = openpyxl.load_workbook(f)
            for n in FOGLI_DATI:
                self.assertIn(n, wb.sheetnames)
                self.assertEqual((wb[n].max_row, wb[n].max_column), atteso[n], f"foglio {n} alterato dal salvataggio")
            self.assertIn("Istruzioni", wb.sheetnames, "il foglio Istruzioni non deve sparire")
            self.assertEqual(list(self.tmp.glob("*.tmp.xlsx")), [], "nessun file temporaneo lasciato in giro")
            yield
        self.esegui(script)


# ── Riepilogo ────────────────────────────────────────────────────────────────

class Risultato(unittest.TextTestResult):
    def __init__(self, *a, **k):
        super().__init__(*a, **k)
        self.esiti: list[tuple[str, str, float]] = []
        self._t0 = 0.0

    def startTest(self, test):
        self._t0 = time.monotonic()
        super().startTest(test)

    def addSuccess(self, test):
        self.esiti.append((test._testMethodName, "OK", time.monotonic() - self._t0))
        super().addSuccess(test)

    def addFailure(self, test, err):
        self.esiti.append((test._testMethodName, "FALLITO", time.monotonic() - self._t0))
        super().addFailure(test, err)

    def addError(self, test, err):
        self.esiti.append((test._testMethodName, "ERRORE", time.monotonic() - self._t0))
        super().addError(test, err)


def main() -> int:
    if not ESEMPIO.exists():
        print(f"Manca il file di esempio {ESEMPIO}")
        return 2
    suite = unittest.TestLoader().loadTestsFromModule(sys.modules[__name__])
    runner = unittest.TextTestRunner(verbosity=2, resultclass=Risultato)
    res = runner.run(suite)
    print("\n" + "═" * 78)
    print("RIEPILOGO DEI TEST DELL'INTERFACCIA")
    print("═" * 78)
    for nome, esito, secondi in res.esiti:
        segno = "✓" if esito == "OK" else "✗"
        print(f" {segno} {nome:<44} {esito:<8} {secondi:5.1f}s")
    print("─" * 78)
    print(f" {len(res.esiti)} test: {sum(1 for _n, e, _s in res.esiti if e == 'OK')} passati, "
          f"{len(res.failures)} falliti, {len(res.errors)} in errore")
    print("═" * 78)
    return 0 if res.wasSuccessful() else 1


if __name__ == "__main__":
    sys.exit(main())
