"""Orario in Word (.docx): stessa impaginazione dei PDF, ma modificabile a mano.

Tre documenti, come per i PDF:
  * `orario_settimanale.docx` — una pagina per giorno, orizzontale, colonne = docenti
  * `orario_docenti.docx`     — una pagina per docente
  * `orario_studenti.docx`    — due ragazzi per pagina, da consegnare alle famiglie
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .costanti import GIORNI, GIORNI_LUNGHI, N_GIORNI, N_ORE
from .export_comune import (
    LEGENDA, LEGENDA_CASELLE, Cella, Riga, cognomi_doppi, etichetta_fascia, etichetta_fascia_breve,
    griglia_docente, griglia_studente, studenti_ordinati, titolo_orario,
)
from .modello import Orario

CARATTERE = "Calibri"
GRIGIO = "E7E6E6"          # celle in cui il docente non è disponibile
AZZURRO = "DDEBF7"         # intestazioni


# ── Mattoni ──────────────────────────────────────────────────────────────────

def _sfondo(cella, colore: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), colore)
    cella._tc.get_or_add_tcPr().append(shd)


def _bordi_tabella(tabella) -> None:
    bordi = OxmlElement("w:tblBorders")
    for lato in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{lato}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "9A9A9A")
        bordi.append(el)
    tabella._tbl.tblPr.append(bordi)


def _ripeti_intestazione(riga) -> None:
    """La riga si ripete in cima a ogni pagina, se la tabella si spezza."""
    pr = riga._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    pr.append(el)


def _scrivi(cella, righe: list[Riga], punti: float, centrato: bool = True,
            grassetto_tutto: bool = False) -> None:
    cella.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    par = cella.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER if centrato else WD_ALIGN_PARAGRAPH.LEFT
    par.paragraph_format.space_before = Pt(1)
    par.paragraph_format.space_after = Pt(1)
    for i, r in enumerate(righe):
        if i:
            par = cella.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER if centrato else WD_ALIGN_PARAGRAPH.LEFT
            par.paragraph_format.space_before = Pt(0)
            par.paragraph_format.space_after = Pt(0)
        run = par.add_run(r.testo)
        run.font.name = CARATTERE
        run.font.size = Pt(punti)
        run.bold = r.grassetto or grassetto_tutto
        run.italic = r.corsivo


def _testo(cella, testo: str, punti: float, grassetto: bool = False, centrato: bool = True) -> None:
    _scrivi(cella, [Riga(testo, grassetto)], punti, centrato)


def _paragrafo(doc, testo: str, punti: float, grassetto: bool = False,
               centrato: bool = True, dopo: float = 6) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centrato else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(dopo)
    run = p.add_run(testo)
    run.font.name = CARATTERE
    run.font.size = Pt(punti)
    run.bold = grassetto


def _pagina(doc, orizzontale: bool, prima: bool = False):
    """Imposta la pagina (A4) o ne apre una nuova con la stessa impostazione."""
    from docx.enum.section import WD_SECTION

    s = doc.sections[0] if prima else doc.add_section(WD_SECTION.NEW_PAGE)
    larghezza, altezza = (Cm(29.7), Cm(21.0)) if orizzontale else (Cm(21.0), Cm(29.7))
    s.orientation = WD_ORIENT.LANDSCAPE if orizzontale else WD_ORIENT.PORTRAIT
    s.page_width, s.page_height = larghezza, altezza
    margine = Cm(0.8) if orizzontale else Cm(1.5)
    s.left_margin = s.right_margin = margine
    s.top_margin = s.bottom_margin = Cm(1.0)
    return s


def _nuovo_documento(orizzontale: bool):
    doc = Document()
    normale = doc.styles["Normal"]
    normale.font.name = CARATTERE
    normale.font.size = Pt(9)
    _pagina(doc, orizzontale, prima=True)
    return doc


def _larghezza_utile(sezione) -> int:
    return sezione.page_width - sezione.left_margin - sezione.right_margin


def _imposta_larghezze(tabella, larghezze: list[int]) -> None:
    tabella.autofit = False
    for riga in tabella.rows:
        for cella, larg in zip(riga.cells, larghezze):
            cella.width = larg


# ── Orario settimanale: una pagina per giorno ────────────────────────────────

def esporta_docx_settimanale(orario: Orario, percorso: Path) -> Path:
    dati = orario.dati
    doppi = cognomi_doppi(dati)
    mappa = orario.per_docente_fascia()
    docenti = dati.docenti
    doc = _nuovo_documento(orizzontale=True)
    sezione = doc.sections[0]
    utile = _larghezza_utile(sezione)
    larg_fascia = Cm(2.2)
    larg_doc = int((utile - larg_fascia) / max(len(docenti), 1))
    punti = 7 if len(docenti) <= 14 else (6 if len(docenti) <= 20 else 5)

    for g in range(N_GIORNI):
        if g:
            _pagina(doc, orizzontale=True)
        _paragrafo(doc, titolo_orario(dati), 14, grassetto=True, dopo=2)
        _paragrafo(doc, GIORNI_LUNGHI[g].upper(), 12, grassetto=True, dopo=6)

        tabella = doc.add_table(rows=3 + N_ORE, cols=1 + len(docenti))
        tabella.alignment = WD_TABLE_ALIGNMENT.CENTER
        _bordi_tabella(tabella)
        intestazioni = [("Docente", lambda d: d.nome, True),
                        ("Aula", lambda d: d.aula(g) or "-", False),
                        ("Strumento", lambda d: d.strumenti.title() or "-", False)]
        for r, (titolo, valore, grassetto) in enumerate(intestazioni):
            _testo(tabella.cell(r, 0), titolo, punti, grassetto=True)
            _sfondo(tabella.cell(r, 0), AZZURRO)
            for c, d in enumerate(docenti, start=1):
                _testo(tabella.cell(r, c), valore(d), punti, grassetto=grassetto)
                _sfondo(tabella.cell(r, c), AZZURRO)
            _ripeti_intestazione(tabella.rows[r])
        for o in range(N_ORE):
            r = 3 + o
            _testo(tabella.cell(r, 0), etichetta_fascia(o), punti, grassetto=True)
            _sfondo(tabella.cell(r, 0), AZZURRO)
            for c, d in enumerate(docenti, start=1):
                from .export_comune import cella as costruisci
                cel = costruisci(orario, d.nome, g * N_ORE + o, doppi, mappa)
                _scrivi(tabella.cell(r, c), cel.righe, punti)
                if cel.non_disponibile:
                    _sfondo(tabella.cell(r, c), GRIGIO)
            tabella.rows[r].height = Cm(2.4)
        _imposta_larghezze(tabella, [larg_fascia] + [larg_doc] * len(docenti))
        _paragrafo(doc, LEGENDA + " — " + LEGENDA_CASELLE, 7, dopo=0)

    if dati.lmi:
        _pagina(doc, orizzontale=True)
        _paragrafo(doc, "LABORATORI DI MUSICA D'INSIEME (ORARIO DEL MATTINO)", 13, grassetto=True, dopo=8)
        colonne = ["Laboratorio", "Classi", "Docente", "Aula", "Giorno e ora", "Studenti"]
        tabella = doc.add_table(rows=1 + len(dati.lmi), cols=len(colonne))
        _bordi_tabella(tabella)
        for c, titolo in enumerate(colonne):
            _testo(tabella.cell(0, c), titolo, 9, grassetto=True)
            _sfondo(tabella.cell(0, c), AZZURRO)
        _ripeti_intestazione(tabella.rows[0])
        for r, l in enumerate(dati.lmi, start=1):
            for c, v in enumerate((l.nome, l.classi, l.docente, l.aula, l.giorno_ora, l.studenti)):
                _testo(tabella.cell(r, c), v or "-", 8, centrato=c < 5)
        _imposta_larghezze(tabella, [Cm(5), Cm(2), Cm(3), Cm(3.5), Cm(3.5), Cm(11)])

    percorso = Path(percorso)
    doc.save(percorso)
    return percorso


# ── Griglia settimanale 4 ore × 5 giorni (docenti e studenti) ────────────────

def _griglia_settimana(doc, griglia: list[list[Cella]], punti: float, altezza: Cm) -> None:
    tabella = doc.add_table(rows=1 + N_ORE, cols=1 + N_GIORNI)
    tabella.alignment = WD_TABLE_ALIGNMENT.CENTER
    _bordi_tabella(tabella)
    _testo(tabella.cell(0, 0), "", punti)
    _sfondo(tabella.cell(0, 0), AZZURRO)
    for c, giorno in enumerate(GIORNI_LUNGHI, start=1):
        _testo(tabella.cell(0, c), giorno, punti, grassetto=True)
        _sfondo(tabella.cell(0, c), AZZURRO)
    _ripeti_intestazione(tabella.rows[0])
    for o in range(N_ORE):
        _testo(tabella.cell(1 + o, 0), etichetta_fascia(o), punti, grassetto=True)
        _sfondo(tabella.cell(1 + o, 0), AZZURRO)
        for g in range(N_GIORNI):
            cel = griglia[o][g]
            _scrivi(tabella.cell(1 + o, 1 + g), cel.righe, punti)
            if cel.non_disponibile:
                _sfondo(tabella.cell(1 + o, 1 + g), GRIGIO)
        tabella.rows[1 + o].height = altezza
    larg = int((_larghezza_utile(doc.sections[-1]) - Cm(2.6)) / N_GIORNI)
    _imposta_larghezze(tabella, [Cm(2.6)] + [larg] * N_GIORNI)


def esporta_docx_docenti(orario: Orario, percorso: Path) -> Path:
    dati = orario.dati
    doppi = cognomi_doppi(dati)
    mappa = orario.per_docente_fascia()
    doc = _nuovo_documento(orizzontale=False)
    for i, d in enumerate(dati.docenti):
        if i:
            _pagina(doc, orizzontale=False)
        lezioni = orario.lezioni_docente(d.nome)
        aule = (f"Aula: {d.aula() or '-'}" if d.aula_unica
                else "Aule: " + ", ".join(f"{GIORNI[k]} {a or '-'}" for k, a in enumerate(d.aule)))
        _paragrafo(doc, titolo_orario(dati), 12, grassetto=True, dopo=2)
        _paragrafo(doc, d.nome, 16, grassetto=True, dopo=2)
        _paragrafo(doc, f"{d.strumenti.title() or '-'}   |   {aule}   |   "
                        f"Ore in orario: {len(lezioni)} (disponibilità: {len(d.disponibilita)} fasce)", 9, dopo=8)
        _griglia_settimana(doc, griglia_docente(orario, d.nome, doppi, mappa), 9, Cm(3.0))
        _paragrafo(doc, LEGENDA + " — " + LEGENDA_CASELLE, 7, dopo=0)
    percorso = Path(percorso)
    doc.save(percorso)
    return percorso


def esporta_docx_studenti(orario: Orario, percorso: Path) -> Path:
    dati = orario.dati
    doppi = cognomi_doppi(dati)
    doc = _nuovo_documento(orizzontale=False)
    studenti = studenti_ordinati(dati)
    for i, s in enumerate(studenti):
        if i and i % 2 == 0:
            _pagina(doc, orizzontale=False)
        elif i:
            _paragrafo(doc, "", 8, dopo=10)
        distanza = (f"{s.trasporto.minuti_min} min di ritorno" if s.trasporto and s.trasporto.minuti_min is not None
                    else (f"{s.km:g} km" if s.km is not None else ""))
        _paragrafo(doc, titolo_orario(dati), 10, grassetto=True, dopo=2)
        _paragrafo(doc, f"{s.cognome.title()} {s.nome.title()}", 14, grassetto=True, dopo=2)
        _paragrafo(doc, "   |   ".join(x for x in (f"Classe {s.classe}ª", distanza) if x), 9, dopo=6)
        _griglia_settimana(doc, griglia_studente(orario, s.id, doppi), 8, Cm(1.9))
    percorso = Path(percorso)
    doc.save(percorso)
    return percorso


__all__ = ["esporta_docx_settimanale", "esporta_docx_docenti", "esporta_docx_studenti"]
