"""Esporta l'orario calcolato in `orario.xlsx` (openpyxl) e coordina l'export completo."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .costanti import GIORNI, GIORNI_LUNGHI, N_GIORNI, N_ORE, ORE
from .export_comune import (
    Cella, cognomi_doppi, etichetta_fascia, griglia_docente, lezione_per_studente_breve,
    riepilogo, studenti_ordinati,
)
from .modello import Orario

# ── Stili ────────────────────────────────────────────────────────────────────

AZZURRO = PatternFill("solid", fgColor="DDEBF7")
GRIGIO = PatternFill("solid", fgColor="E7E6E6")
GIALLO = PatternFill("solid", fgColor="FFF2CC")
_lato = Side(style="thin", color="999999")
BORDO = Border(left=_lato, right=_lato, top=_lato, bottom=_lato)
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
SINISTRA = Alignment(horizontal="left", vertical="top", wrap_text=True)
INTESTAZIONE = Font(bold=True, size=10)
LARGHEZZA_DOCENTE = 16
LARGHEZZA_FASCIA = 15


def _intestazione(c, testo: str, larghezza: int | None = None, ws: Worksheet | None = None) -> None:
    c.value = testo
    c.font = INTESTAZIONE
    c.fill = AZZURRO
    c.alignment = CENTRO
    c.border = BORDO
    if larghezza and ws is not None:
        ws.column_dimensions[c.column_letter].width = larghezza


def _scrivi_cella(c, cella: Cella, allinea=CENTRO) -> None:
    c.border = BORDO
    c.alignment = allinea
    if cella.vuota:
        if cella.non_disponibile:
            c.fill = GRIGIO
        return
    c.value = cella.testo
    c.font = Font(bold=cella.grassetto, italic=cella.corsivo, size=10)


def _colonne_fasce(ws: Worksheet) -> None:
    ws.column_dimensions["A"].width = LARGHEZZA_FASCIA


# ── Fogli per giorno ─────────────────────────────────────────────────────────

def _foglio_giorno(wb: Workbook, orario: Orario, g: int, doppi: set[str], mappa) -> None:
    ws = wb.create_sheet(GIORNI[g])
    dati = orario.dati
    _colonne_fasce(ws)
    _intestazione(ws.cell(1, 1), GIORNI_LUNGHI[g])
    _intestazione(ws.cell(2, 1), "Aula")
    _intestazione(ws.cell(3, 1), "Strumento")
    for j, d in enumerate(dati.docenti, start=2):
        _intestazione(ws.cell(1, j), d.nome, LARGHEZZA_DOCENTE, ws)
        _intestazione(ws.cell(2, j), d.aula(g) or "-")
        _intestazione(ws.cell(3, j), d.strumenti.title() or "-")
    for o in range(N_ORE):
        r = 4 + o
        _intestazione(ws.cell(r, 1), etichetta_fascia(o))
        ws.row_dimensions[r].height = 64
    for j, d in enumerate(dati.docenti, start=2):
        griglia = griglia_docente(orario, d.nome, doppi, mappa)
        for o in range(N_ORE):
            _scrivi_cella(ws.cell(4 + o, j), griglia[o][g])
    ws.freeze_panes = "B4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:3"


# ── Foglio Studenti ──────────────────────────────────────────────────────────

def _foglio_studenti(wb: Workbook, orario: Orario) -> None:
    ws = wb.create_sheet("Studenti")
    colonne = ["Classe", "Cognome", "Nome", "KM"] + GIORNI
    larghezze = [7, 18, 16, 6] + [26] * N_GIORNI
    for j, (nome, larg) in enumerate(zip(colonne, larghezze), start=1):
        _intestazione(ws.cell(1, j), nome, larg, ws)
    r = 2
    for s in studenti_ordinati(orario.dati):
        lezioni = orario.lezioni_di(s.id)
        valori = [s.classe, s.cognome.title(), s.nome.title(), s.km]
        for j, v in enumerate(valori, start=1):
            c = ws.cell(r, j, v)
            c.border = BORDO
            c.alignment = Alignment(horizontal="center" if j not in (2, 3) else "left", vertical="top")
        if not lezioni:
            ws.cell(r, 2).fill = GIALLO
        for g in range(N_GIORNI):
            testi = [lezione_per_studente_breve(orario, l) for l in lezioni if l.giorno == g]
            c = ws.cell(r, 5 + g, "\n".join(testi) or None)
            c.border = BORDO
            c.alignment = SINISTRA
        ws.row_dimensions[r].height = max(15, 14 * max((sum(1 for l in lezioni if l.giorno == g) for g in range(N_GIORNI)), default=1))
        r += 1
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(colonne))}{max(r - 1, 1)}"


# ── Foglio Docenti ───────────────────────────────────────────────────────────

def _foglio_docenti(wb: Workbook, orario: Orario, doppi: set[str], mappa) -> None:
    ws = wb.create_sheet("Docenti")
    dati = orario.dati
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 14
    _intestazione(ws.cell(1, 1), "Giorno")
    _intestazione(ws.cell(2, 1), "")
    _intestazione(ws.cell(3, 1), "")
    _intestazione(ws.cell(1, 2), "Fascia")
    _intestazione(ws.cell(2, 2), "Aula")
    _intestazione(ws.cell(3, 2), "Strumento")
    for j, d in enumerate(dati.docenti, start=3):
        _intestazione(ws.cell(1, j), d.nome, LARGHEZZA_DOCENTE, ws)
        _intestazione(ws.cell(2, j), d.aula() or "varia")   # nel riepilogo l'aula cambia di giorno in giorno
        _intestazione(ws.cell(3, j), d.strumenti.title() or "-")
    griglie = {d.nome: griglia_docente(orario, d.nome, doppi, mappa) for d in dati.docenti}
    for g in range(N_GIORNI):
        r0 = 4 + g * N_ORE
        for j, d in enumerate(dati.docenti, start=3):
            if not d.aula_unica and d.aula(g):
                _intestazione(ws.cell(r0, j), f"aula {d.aula(g)}")
        ws.merge_cells(start_row=r0, start_column=1, end_row=r0 + N_ORE - 1, end_column=1)
        _intestazione(ws.cell(r0, 1), GIORNI_LUNGHI[g])
        for o in range(N_ORE):
            r = r0 + o
            ws.cell(r, 1).border = BORDO
            _intestazione(ws.cell(r, 2), etichetta_fascia(o))
            ws.row_dimensions[r].height = 48
            for j, d in enumerate(dati.docenti, start=3):
                _scrivi_cella(ws.cell(r, j), griglie[d.nome][o][g])
    ws.freeze_panes = "C4"


# ── Foglio Controlli ─────────────────────────────────────────────────────────

def _foglio_controlli(wb: Workbook, orario: Orario) -> None:
    ws = wb.create_sheet("Controlli")
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 48
    ws.column_dimensions["C"].width = 90
    ws.cell(1, 1, "Riepilogo dell'orario").font = Font(bold=True, size=12)
    r = 2
    for voce, valore in riepilogo(orario):
        a = ws.cell(r, 1, voce)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        a.border = BORDO
        a.alignment = SINISTRA
        v = ws.cell(r, 3, valore)
        v.border = BORDO
        v.alignment = Alignment(horizontal="left", vertical="top")
        r += 1
    r += 1
    ws.cell(r, 1, "Avvisi: preferenze non soddisfatte e forzature").font = Font(bold=True, size=12)
    r += 1
    for j, nome in enumerate(["Livello", "Dove", "Messaggio"], start=1):
        _intestazione(ws.cell(r, j), nome)
    r += 1
    if not orario.avvisi:
        ws.cell(r, 1, "Nessun avviso: tutte le preferenze sono state rispettate.").alignment = SINISTRA
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        return
    for p in orario.avvisi:
        for j, v in enumerate([p.livello, p.dove, p.messaggio], start=1):
            c = ws.cell(r, j, v)
            c.border = BORDO
            c.alignment = SINISTRA
        if p.livello == "errore":
            ws.cell(r, 1).fill = GIALLO
        r += 1


# ── Foglio LMI ───────────────────────────────────────────────────────────────

def _foglio_lmi(wb: Workbook, orario: Orario) -> None:
    ws = wb.create_sheet("LMI")
    colonne = ["Laboratorio", "Classi", "Docente", "Aula", "Giorno e ora", "Studenti", "Note"]
    larghezze = [24, 10, 16, 22, 20, 50, 30]
    for j, (nome, larg) in enumerate(zip(colonne, larghezze), start=1):
        _intestazione(ws.cell(1, j), nome, larg, ws)
    if not orario.dati.lmi:
        ws.cell(2, 1, "Nessun laboratorio indicato nel file di input").alignment = SINISTRA
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(colonne))
        return
    for r, l in enumerate(orario.dati.lmi, start=2):
        for j, v in enumerate([l.nome, l.classi, l.docente, l.aula, l.giorno_ora, l.studenti, l.note], start=1):
            c = ws.cell(r, j, v or None)
            c.border = BORDO
            c.alignment = SINISTRA
    ws.freeze_panes = "A2"


# ── API pubblica ─────────────────────────────────────────────────────────────

def esporta_excel(orario: Orario, percorso: Path) -> Path:
    percorso = Path(percorso)
    percorso.parent.mkdir(parents=True, exist_ok=True)
    doppi = cognomi_doppi(orario.dati)
    mappa = orario.per_docente_fascia()
    wb = Workbook()
    wb.remove(wb.active)
    for g in range(N_GIORNI):
        _foglio_giorno(wb, orario, g, doppi, mappa)
    _foglio_studenti(wb, orario)
    _foglio_docenti(wb, orario, doppi, mappa)
    _foglio_controlli(wb, orario)
    _foglio_lmi(wb, orario)
    wb.save(percorso)
    return percorso


def esporta_tutto(orario: Orario, cartella: Path) -> list[Path]:
    """Produce Excel, i tre PDF e i tre Word nella cartella indicata (creata se manca)."""
    from .export_docx import esporta_docx_docenti, esporta_docx_settimanale, esporta_docx_studenti
    from .export_pdf import esporta_pdf_docenti, esporta_pdf_settimanale, esporta_pdf_studenti

    cartella = Path(cartella)
    cartella.mkdir(parents=True, exist_ok=True)
    return [
        esporta_excel(orario, cartella / "orario.xlsx"),
        esporta_pdf_settimanale(orario, cartella / "orario_settimanale.pdf"),
        esporta_pdf_docenti(orario, cartella / "orario_docenti.pdf"),
        esporta_pdf_studenti(orario, cartella / "orario_studenti.pdf"),
        esporta_docx_settimanale(orario, cartella / "orario_settimanale.docx"),
        esporta_docx_docenti(orario, cartella / "orario_docenti.docx"),
        esporta_docx_studenti(orario, cartella / "orario_studenti.docx"),
    ]


__all__ = ["esporta_excel", "esporta_tutto"]
