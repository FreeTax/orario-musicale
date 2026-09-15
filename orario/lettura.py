"""Lettura e scrittura del file di input `input_orario.xlsx`.

Due livelli:
  * tabelle "grezze" (intestazione + righe di stringhe) → usate dalla finestra per la griglia
    modificabile e per il salvataggio senza perdere stili e menu a tendina;
  * `DatiInput` strutturato → costruito dalle tabelle grezze, usato da controlli e motore.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from datetime import datetime

from .costanti import (
    COLONNE_AULE, FASCE, FOGLI_DATI, FOGLI_FACOLTATIVI, FOGLIO_ABBINAMENTI, FOGLIO_DOCENTI, FOGLIO_GRUPPI,
    FOGLIO_IMPEGNI, FOGLIO_LMI, FOGLIO_ORARIO, FOGLIO_PARAMETRI, FOGLIO_STUDENTI, FOGLIO_TRASPORTI, GIORNI,
    GIORNI_LUNGHI, NOMI_TIPO, ORE, ORE_FINE, fascia, tipo_da_testo,
)
from .modello import (
    Abbinamento, DatiInput, Docente, GruppoLMC, LaboratorioLMI, Lezione, Orario, Parametri, Problema,
    ProblemiError, Studente, Trasporto,
)


@dataclass
class Tabella:
    nome: str
    intestazione: list[str]
    righe: list[list[str]] = field(default_factory=list)

    def colonna(self, prefisso: str) -> int:
        """Indice della prima colonna la cui intestazione inizia con `prefisso` (case-insensitive)."""
        p = prefisso.lower()
        for i, h in enumerate(self.intestazione):
            if (h or "").strip().lower().startswith(p):
                return i
        raise KeyError(f"Colonna '{prefisso}' non trovata nel foglio '{self.nome}'")

    def valore(self, riga: list[str], prefisso: str) -> str:
        try:
            i = self.colonna(prefisso)
        except KeyError:
            return ""
        return riga[i].strip() if i < len(riga) and riga[i] is not None else ""


def _testo(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


# ── Tabelle grezze ───────────────────────────────────────────────────────────

def leggi_tabelle(percorso: Path) -> dict[str, Tabella]:
    """Legge i fogli dati come tabelle di stringhe. Solleva ProblemiError se manca un foglio."""
    wb = load_workbook(percorso, data_only=True)
    mancanti = [f for f in FOGLI_DATI if f not in wb.sheetnames and f not in FOGLI_FACOLTATIVI]
    if mancanti:
        raise ProblemiError([Problema(
            "errore", "File",
            f"Mancano i fogli: {', '.join(mancanti)}. Il file deve avere i fogli {', '.join(FOGLI_DATI)}.",
        )])
    from .template import COLONNE_ABBINAMENTI, COLONNE_IMPEGNI

    tabelle: dict[str, Tabella] = {}
    for nome in FOGLI_DATI:
        if nome not in wb.sheetnames:      # foglio aggiunto in una versione successiva
            vuoti = {FOGLIO_IMPEGNI: COLONNE_IMPEGNI, FOGLIO_ABBINAMENTI: COLONNE_ABBINAMENTI}
            tabelle[nome] = Tabella(nome, list(vuoti.get(nome, [])))
            continue
        ws = wb[nome]
        righe = list(ws.iter_rows(values_only=True))
        if not righe:
            tabelle[nome] = Tabella(nome, [])
            continue
        intestazione = [_testo(c) for c in righe[0]]
        # taglia le colonne vuote in coda all'intestazione
        while intestazione and not intestazione[-1]:
            intestazione.pop()
        n = len(intestazione)
        dati = []
        for r in righe[1:]:
            valori = [_testo(c) for c in r[:n]] + [""] * max(0, n - len(r))
            if any(valori):
                dati.append(valori)
        tabelle[nome] = Tabella(nome, intestazione, dati)
    return tabelle


def salva_tabelle(percorso: Path, tabelle: dict[str, Tabella]) -> None:
    """Riscrive i valori dei fogli dati nel file esistente, preservando stili e menu a tendina."""
    wb = load_workbook(percorso)
    for nome, tab in tabelle.items():
        if nome not in wb.sheetnames:
            if not tab.intestazione:
                continue
            _crea_foglio_mancante(wb, nome, tab)
        ws = wb[nome]
        # svuota i valori delle righe dati (non le formattazioni)
        for row in ws.iter_rows(min_row=2, max_row=max(ws.max_row, 2)):
            for c in row:
                c.value = None
        for i, h in enumerate(tab.intestazione, start=1):
            ws.cell(row=1, column=i).value = h
        for r, riga in enumerate(tab.righe, start=2):
            for c, v in enumerate(riga, start=1):
                ws.cell(row=r, column=c).value = _numero_o_testo(v)
    _salva_atomico(wb, percorso)


def _salva_atomico(wb, percorso: Path) -> None:
    """Scrive su un file temporaneo e poi lo rinomina, con qualche tentativo.

    Su Windows la rinomina può fallire per un istante se il file è tenuto aperto da antivirus o Excel.
    """
    import time

    tmp = percorso.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    ultimo: Exception | None = None
    for tentativo in range(5):
        try:
            tmp.replace(percorso)
            return
        except PermissionError as e:
            ultimo = e
            time.sleep(0.4 * (tentativo + 1))
    tmp.unlink(missing_ok=True)
    raise PermissionError(f"Non riesco a sovrascrivere {percorso.name}: è aperto in un altro programma "
                          f"(Excel?). Chiuderlo e riprovare.") from ultimo


def _crea_foglio_mancante(wb, nome: str, tab: Tabella) -> None:
    """Aggiunge al file un foglio introdotto in una versione successiva, con la stessa impaginazione."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    posizione = FOGLI_DATI.index(nome) if nome in FOGLI_DATI else None
    ws = wb.create_sheet(nome, posizione + 1 if posizione is not None else None)
    ws.append(list(tab.intestazione))
    for i, titolo in enumerate(tab.intestazione, start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
        stretta = titolo in FASCE
        c.alignment = Alignment(wrap_text=True, vertical="bottom" if stretta else "center",
                                text_rotation=90 if stretta else 0,
                                horizontal="center" if stretta else "general")
        ws.column_dimensions[get_column_letter(i)].width = 6.5 if stretta else (22 if i in (2, 3) else 14)
    ws.row_dimensions[1].height = 60 if any(t in FASCE for t in tab.intestazione) else 32
    ws.freeze_panes = "A2"


def _numero_o_testo(v: str):
    if v is None or v == "":
        return None
    s = str(v).strip()
    try:
        if re.fullmatch(r"-?\d+", s):
            return int(s)
        if re.fullmatch(r"-?\d+[.,]\d+", s):
            return float(s.replace(",", "."))
    except ValueError:
        pass
    return s


# ── Righe di esempio del modello ─────────────────────────────────────────────

# I nomi inventati che il modello mette nelle righe grigie di esempio.
_ESEMPIO_STUDENTI = {("ROSSI", "MARIO"), ("NERI", "ANNA")}
_ESEMPIO_DOCENTI = {"BIANCHI", "VERDI"}
_ESEMPIO_GRUPPI = {"NERI", "ROSSI"}


def _nota_di_esempio(nota: str) -> bool:
    return "ESEMPIO" in (nota or "").upper()


def _residuo_esempio(dove: str, riga: int) -> Problema:
    return Problema("avviso", f"{dove}, riga {riga}",
                    "Nella colonna Note è rimasta la scritta «ESEMPIO»: la riga contiene dati veri, "
                    "quindi è stata usata lo stesso. Puoi cancellare quella nota.",
                    "righe con la nota ESEMPIO rimasta")


# ── Costruzione DatiInput ────────────────────────────────────────────────────

_SI = {"SI", "SÌ", "S", "YES", "X", "1", "TRUE", "VERO"}


def _bool(v: str) -> bool:
    return (v or "").strip().upper() in _SI


def _si_no(v: str) -> bool | None:
    """SI → True, NO → False, casella vuota → None (nessuna preferenza)."""
    t = (v or "").strip().upper()
    if not t:
        return None
    return t in _SI


def _float(v: str) -> float | None:
    s = (v or "").strip().replace(",", ".")
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _int(v: str, default: int = 0) -> int:
    s = (v or "").strip()
    if not s:
        return default
    try:
        return int(float(s.replace(",", ".")))
    except ValueError:
        return default


def _giorni(v: str) -> tuple[set[int], list[str]]:
    """'Mar, Gio' → {1, 3}. Ritorna anche i pezzi non riconosciuti."""
    out, sconosciuti = set(), []
    for pezzo in re.split(r"[,;/ ]+", (v or "").strip()):
        if not pezzo:
            continue
        p = pezzo.lower().rstrip(".")
        trovato = False
        for i, (g, gl) in enumerate(zip(GIORNI, GIORNI_LUNGHI)):
            if p == g.lower() or p == gl.lower() or gl.lower().startswith(p[:3]):
                out.add(i)
                trovato = True
                break
        if not trovato:
            sconosciuti.append(pezzo)
    return out, sconosciuti


def costruisci_dati(tabelle: dict[str, Tabella], percorso: Path | None = None) -> DatiInput:
    """Trasforma le tabelle grezze in DatiInput. Errori di formato → ProblemiError.

    I controlli di coerenza (nomi che non combaciano, ore, ecc.) sono in `controlli.py`;
    qui si segnalano solo i valori che non si riescono proprio a leggere.
    """
    problemi: list[Problema] = []

    # Studenti
    ts = tabelle[FOGLIO_STUDENTI]
    studenti: list[Studente] = []
    for i, r in enumerate(ts.righe, start=2):
        v = lambda p: ts.valore(r, p)  # noqa: E731
        cognome, nome = v("Cognome"), v("Nome")
        if _nota_di_esempio(v("Note")):
            if (cognome.upper(), nome.upper()) in _ESEMPIO_STUDENTI:
                problemi.append(Problema("errore", f"{FOGLIO_STUDENTI}, riga {i}",
                                         "Riga di esempio ancora presente: cancellarla o sostituirla con uno studente vero."))
                continue
            problemi.append(_residuo_esempio(FOGLIO_STUDENTI, i))
        if not cognome:
            problemi.append(Problema("errore", f"{FOGLIO_STUDENTI}, riga {i}", "Cognome mancante."))
            continue
        classe = _int(v("Classe"), 0)
        if classe not in (1, 2, 3, 4, 5):
            problemi.append(Problema("errore", f"{FOGLIO_STUDENTI}, riga {i} ({cognome} {nome})",
                                     f"Classe '{v('Classe')}' non valida: deve essere 1, 2, 3, 4 o 5."))
            continue
        km_txt = v("KM")
        km = _float(km_txt)
        if km_txt and km is None:
            problemi.append(Problema("errore", f"{FOGLIO_STUDENTI}, riga {i} ({cognome} {nome})",
                                     f"KM '{km_txt}' non è un numero."))
        giorni, sconosciuti = _giorni(v("Giorni NON"))
        if sconosciuti:
            problemi.append(Problema("errore", f"{FOGLIO_STUDENTI}, riga {i} ({cognome} {nome})",
                                     f"Giorni non riconosciuti: {', '.join(sconosciuti)}. Usare Lun, Mar, Mer, Gio, Ven."))
        studenti.append(Studente(
            classe=classe, cognome=cognome.upper(), nome=nome.upper(), km=km,
            strum1=v("Strumento 1").upper(), doc1=v("Docente 1"),
            strum2=v("Strumento 2").upper() if classe != 5 else "",
            doc2=v("Docente 2") if classe != 5 else "",
            # il nome vecchio della colonna era "Ore consecutive": i file di prima si leggono lo stesso
            primo_attaccato=_si_no(v("1° strumento attaccato") or v("Ore consecutive")),
            giorno_unico=_bool(v("Giorno unico")),
            giorni_non_disp=giorni, note=v("Note"), riga=i,
            comune=v("Comune"), indirizzo=v("Indirizzo"), civico=v("Civico"),
            minuti_manuali=_float(v("Minuti per tornare a casa")),
        ))

    # Docenti
    td = tabelle[FOGLIO_DOCENTI]
    docenti: list[Docente] = []
    idx_fasce: list[int | None] = []
    for nome_colonna in FASCE:   # non chiamarla "fascia": nasconderebbe la funzione omonima
        try:
            idx_fasce.append(td.colonna(nome_colonna))
        except KeyError:
            idx_fasce.append(None)
    if any(i is None for i in idx_fasce):
        mancanti = [f for f, i in zip(FASCE, idx_fasce) if i is None]
        problemi.append(Problema("errore", FOGLIO_DOCENTI,
                                 f"Mancano le colonne di disponibilità: {', '.join(mancanti)}."))
    for i, r in enumerate(td.righe, start=2):
        v = lambda p: td.valore(r, p)  # noqa: E731
        nome = v("Docente")
        if not nome:
            continue
        if nome.strip().upper() == "TOTALE":   # riga di comodo con le somme
            continue
        if _nota_di_esempio(v("Note")):
            if nome.upper() in _ESEMPIO_DOCENTI:
                problemi.append(Problema("errore", f"{FOGLIO_DOCENTI}, riga {i}",
                                         "Riga di esempio ancora presente: cancellarla o sostituirla con un docente vero."))
                continue
            problemi.append(_residuo_esempio(FOGLIO_DOCENTI, i))
        disp: dict[int, str] = {}
        for f, col in enumerate(idx_fasce):
            if col is None or col >= len(r):
                continue
            cella = (r[col] or "").strip().upper()
            if cella in ("X", "A"):
                disp[f] = cella
            elif cella:
                problemi.append(Problema("errore", f"{FOGLIO_DOCENTI}: {nome}",
                                         f"Valore '{r[col]}' nella colonna {FASCE[f]}: ammessi solo X, A o vuoto."))
        accomp_txt = v("Ore accomp")
        accomp = _int(accomp_txt, 0)
        if accomp_txt and not re.fullmatch(r"\d+", accomp_txt):
            problemi.append(Problema("errore", f"{FOGLIO_DOCENTI}: {nome}",
                                     f"Ore accompagnamento '{accomp_txt}' non è un numero intero."))
        aule = [v(nome_colonna) for nome_colonna in COLONNE_AULE]
        if not any(aule):  # file creato quando l'aula era una sola per tutta la settimana
            unica = v("Aula")
            aule = [unica] * len(COLONNE_AULE)
        docenti.append(Docente(nome=nome, strumenti=v("Strumento"), aule=aule,
                               ore_accomp=accomp, disponibilita=disp, note=v("Note"), riga=i))

    # Gruppi LMC (gli id studente vengono risolti dai controlli)
    tg = tabelle[FOGLIO_GRUPPI]
    gruppi: list[GruppoLMC] = []
    for i, r in enumerate(tg.righe, start=2):
        v = lambda p: tg.valore(r, p)  # noqa: E731
        membri = [v(f"Studente {k}") for k in range(1, 6)]
        membri = [m for m in membri if m]
        docente = v("Docente")
        if not membri and not docente:
            continue
        if _nota_di_esempio(v("Note")):
            if {m.split()[0].upper() for m in membri} <= _ESEMPIO_GRUPPI:
                problemi.append(Problema("errore", f"{FOGLIO_GRUPPI}, riga {i}",
                                         "Riga di esempio ancora presente: cancellarla o sostituirla con un gruppo vero."))
                continue
            problemi.append(_residuo_esempio(FOGLIO_GRUPPI, i))
        numero = _int(v("Gruppo"), i - 1)
        gruppi.append(GruppoLMC(numero=numero, docente=docente, studenti=[], studenti_raw=membri,
                                note=v("Note"), riga=i))

    # LMI
    tl = tabelle[FOGLIO_LMI]
    lmi: list[LaboratorioLMI] = []
    for i, r in enumerate(tl.righe, start=2):
        v = lambda p: tl.valore(r, p)  # noqa: E731
        if not any(r):
            continue
        if _nota_di_esempio(v("Note")) and v("Laboratorio").strip().upper() == "LMI ARCHI":
            continue
        lmi.append(LaboratorioLMI(nome=v("Laboratorio"), classi=v("Classi"), docente=v("Docente"),
                                  aula=v("Aula"), giorno_ora=v("Giorno"), studenti=v("Studenti"),
                                  note=v("Note"), riga=i))

    # Impegni degli studenti: la casella segnata è un'ora in cui il ragazzo NON c'è
    ti = tabelle[FOGLIO_IMPEGNI]
    if ti.intestazione:
        per_chiave: dict[tuple[str, str], Studente] = {}
        for st in studenti:
            per_chiave[(st.cognome.upper(), st.nome.upper())] = st
        idx_i: list[int | None] = []
        for nome_fascia_col in FASCE:
            try:
                idx_i.append(ti.colonna(nome_fascia_col))
            except KeyError:
                idx_i.append(None)
        for i, r in enumerate(ti.righe, start=2):
            cognome = ti.valore(r, "Cognome").upper()
            nome = ti.valore(r, "Nome").upper()
            if not cognome:
                continue
            st = per_chiave.get((cognome, nome))
            if st is None:
                candidati = [x for k, x in per_chiave.items() if k[0] == cognome]
                st = candidati[0] if len(candidati) == 1 else None
            occupate = {f for f, col in enumerate(idx_i)
                        if col is not None and col < len(r) and (r[col] or "").strip()}
            if st is None:
                if occupate:
                    problemi.append(Problema("avviso", f"{FOGLIO_IMPEGNI}, riga {i}",
                                             f"«{ti.valore(r, 'Cognome')} {ti.valore(r, 'Nome')}» non è nel foglio "
                                             f"{FOGLIO_STUDENTI}: gli impegni segnati su questa riga sono stati "
                                             "ignorati. Se il ragazzo non c'è più, la riga si può cancellare.",
                                             "impegni riferiti a ragazzi non in elenco"))
                continue
            st.fasce_non_disp |= occupate

    # Abbinamenti fissi: lezioni già decise
    ta = tabelle[FOGLIO_ABBINAMENTI]
    abbinamenti: list[Abbinamento] = []
    for i, r in enumerate(ta.righe, start=2):
        v = lambda p: ta.valore(r, p)  # noqa: E731
        docente_nome, studente_txt = v("Docente"), v("Studente")
        if not docente_nome and not studente_txt:
            continue
        if "ESEMPIO" in v("Note").upper() and studente_txt.upper().startswith(("NERI", "ROSSI")):
            continue
        giorno_txt, ora_txt = v("Giorno"), v("Ora")
        giorni_letti, _ = _giorni(giorno_txt)
        if len(giorni_letti) != 1:
            problemi.append(Problema("errore", f"{FOGLIO_ABBINAMENTI}, riga {i}",
                                     f"Giorno «{giorno_txt}» non riconosciuto: scrivere Lun, Mar, Mer, Gio o Ven."))
            continue
        ora_pulita = ora_txt.strip().replace(".", ":")
        if ora_pulita not in ORE:
            problemi.append(Problema("errore", f"{FOGLIO_ABBINAMENTI}, riga {i}",
                                     f"Ora «{ora_txt}» non riconosciuta: sono ammesse {', '.join(ORE)}."))
            continue
        tipo_txt = v("Tipo")
        tipo = tipo_da_testo(tipo_txt)
        if tipo_txt.strip() and not tipo:
            problemi.append(Problema("errore", f"{FOGLIO_ABBINAMENTI}, riga {i}",
                                     f"Tipo di lezione «{tipo_txt}» non riconosciuto: "
                                     f"sono ammessi {', '.join(NOMI_TIPO)} oppure la casella vuota."))
            continue
        abbinamenti.append(Abbinamento(
            docente=docente_nome, studente_raw=studente_txt, tipo=tipo,
            fascia=fascia(giorni_letti.pop(), ORE.index(ora_pulita)), note=v("Note"), riga=i))

    # Parametri
    tp = tabelle[FOGLIO_PARAMETRI]
    par = Parametri()
    for r in tp.righe:
        chiave = tp.valore(r, "Parametro").lower()
        val = tp.valore(r, "Valore")
        if not chiave:
            continue
        if "indirizzo" in chiave and "scuola" in chiave:
            par.indirizzo_scuola = val
        elif "vicino" in chiave and "soglia" in chiave and "minut" in chiave:
            par.soglia_min_vicino = _int(val, par.soglia_min_vicino)
        elif "vicino" in chiave and "soglia" in chiave:
            par.soglia_km_vicino = _float(val) or par.soglia_km_vicino
        elif "vicino" in chiave:
            par.max_rientri_vicini = _int(val, par.max_rientri_vicini)
        elif "rientri" in chiave:
            par.max_rientri = _int(val, par.max_rientri)
        elif "tempo" in chiave or "second" in chiave:
            par.timeout_s = _int(val, par.timeout_s)

    if any(p.livello == "errore" for p in problemi):
        raise ProblemiError(problemi)
    return DatiInput(studenti=studenti, docenti=docenti, gruppi=gruppi, lmi=lmi, parametri=par,
                     abbinamenti=abbinamenti, percorso=percorso, avvisi=problemi)


def leggi_dati(percorso: Path) -> DatiInput:
    """Scorciatoia: file → DatiInput con i trasporti agganciati (senza controlli di coerenza)."""
    dati = costruisci_dati(leggi_tabelle(percorso), percorso)
    applica_trasporti(dati, leggi_trasporti(percorso))
    return dati


# ── Orario calcolato, salvato dentro il file di input ────────────────────────

def salva_orario_nel_file(percorso: Path, orario: Orario) -> None:
    """Scrive (o sostituisce) il foglio 'Orario calcolato' nel file di input.

    Serve al ricalcolo: la volta dopo il programma parte da qui e sposta solo il necessario.
    """
    from .motore import VERSIONE_REGOLE

    wb = load_workbook(percorso)
    if FOGLIO_ORARIO in wb.sheetnames:
        del wb[FOGLIO_ORARIO]
    ws = wb.create_sheet(FOGLIO_ORARIO)
    ws.append(["Docente", "Giorno", "Ora", "Tipo", "Studenti", "Gruppo", "Seconda ora",
               f"Calcolato il {datetime.now():%d/%m/%Y %H:%M}", "Non modificare a mano: lo scrive il programma.",
               f"Regole v{VERSIONE_REGOLE}"])
    for l in sorted(orario.lezioni, key=lambda l: (l.docente, l.fascia)):
        ws.append([l.docente, GIORNI[l.giorno], ORE[l.ora], l.tipo, "; ".join(l.studenti),
                   l.gruppo if l.gruppo is not None else "", "SI" if l.due_ore else "NO"])
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["E"].width = 60
    ws.column_dimensions["H"].width = 28
    _salva_atomico(wb, percorso)


def versione_orario_precedente(percorso: Path) -> int:
    """Versione delle regole con cui è stato calcolato l'orario salvato nel file (0 se non c'è o è vecchio)."""
    try:
        wb = load_workbook(percorso, read_only=True, data_only=True)
        if FOGLIO_ORARIO not in wb.sheetnames:
            return 0
        prima = next(wb[FOGLIO_ORARIO].iter_rows(min_row=1, max_row=1, values_only=True), ())
    except Exception:
        return 0
    for c in prima:
        if isinstance(c, str) and c.startswith("Regole v") and c[8:].isdigit():
            return int(c[8:])
    return 0


def leggi_orario_precedente(percorso: Path) -> tuple[list[Lezione], str] | None:
    """Legge il foglio 'Orario calcolato' se esiste. Ritorna (lezioni, 'data del calcolo') oppure None."""
    wb = load_workbook(percorso, data_only=True)
    if FOGLIO_ORARIO not in wb.sheetnames:
        return None
    ws = wb[FOGLIO_ORARIO]
    righe = list(ws.iter_rows(values_only=True))
    if len(righe) < 2:
        return None
    quando = ""
    for c in righe[0]:
        if isinstance(c, str) and c.startswith("Calcolato il "):
            quando = c[len("Calcolato il "):]
    lezioni: list[Lezione] = []
    for r in righe[1:]:
        if not r or not r[0]:
            continue
        docente, giorno, ora, tipo, studenti, gruppo, due = (list(r) + [None] * 7)[:7]
        try:
            g = GIORNI.index(str(giorno).strip()[:3].title())
            o = ORE.index(str(ora).strip()[:5])
        except ValueError:
            continue
        ids = [x.strip() for x in str(studenti or "").split(";") if x.strip()]
        lezioni.append(Lezione(
            docente=str(docente).strip(), fascia=fascia(g, o), tipo=str(tipo).strip(), studenti=ids,
            gruppo=int(gruppo) if gruppo not in (None, "") else None,
            due_ore=str(due or "").strip().upper() == "SI",
        ))
    return lezioni, quando


# ── Trasporti (ritorno a casa con i mezzi), salvati dentro il file di input ─────

def salva_trasporti(percorso: Path, trasporti: dict[str, Trasporto]) -> None:
    wb = load_workbook(percorso)
    if FOGLIO_TRASPORTI in wb.sheetnames:
        del wb[FOGLIO_TRASPORTI]
    ws = wb.create_sheet(FOGLIO_TRASPORTI)
    testa = ["Studente", "Indirizzo usato", "Lat", "Lon", "Esito", "Aggiornato il"]
    for fine in ORE_FINE:
        testa += [f"Minuti se finisce alle {fine}", f"A casa alle ({fine})", f"Mezzi ({fine})"]
    testa.append("Scritto dal programma (menu Orario → Aggiorna trasporti): non modificare a mano.")
    ws.append(testa)
    for id_, t in sorted(trasporti.items()):
        riga = [id_, t.indirizzo_usato, t.lat, t.lon, t.esito, t.aggiornato]
        for o in range(4):
            riga += [t.minuti[o] if o < len(t.minuti) else None,
                     t.arrivi[o] if o < len(t.arrivi) else "", t.mezzi[o] if o < len(t.mezzi) else ""]
        ws.append(riga)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["E"].width = 22
    ws.freeze_panes = "B2"
    _salva_atomico(wb, percorso)


def leggi_trasporti(percorso: Path) -> dict[str, Trasporto]:
    wb = load_workbook(percorso, data_only=True)
    if FOGLIO_TRASPORTI not in wb.sheetnames:
        return {}
    out: dict[str, Trasporto] = {}
    for r in wb[FOGLIO_TRASPORTI].iter_rows(min_row=2, values_only=True):
        if not r or not r[0]:
            continue
        r = list(r) + [None] * 20
        minuti, arrivi, mezzi = [], [], []
        for o in range(4):
            m = r[6 + 3 * o]
            minuti.append(int(m) if isinstance(m, (int, float)) else None)
            arrivi.append(_testo(r[7 + 3 * o]))
            mezzi.append(_testo(r[8 + 3 * o]))
        out[str(r[0]).strip()] = Trasporto(
            minuti, arrivi, mezzi,
            float(r[2]) if isinstance(r[2], (int, float)) else None,
            float(r[3]) if isinstance(r[3], (int, float)) else None,
            _testo(r[1]), _testo(r[4]), _testo(r[5]),
        )
    return out


def applica_trasporti(dati: DatiInput, trasporti: dict[str, Trasporto]) -> int:
    """Aggancia i trasporti agli studenti (solo se l'indirizzo non è cambiato). Ritorna quanti agganciati."""
    n = 0
    for s in dati.studenti:
        t = trasporti.get(s.id)
        if t and t.valido and (not t.indirizzo_usato or t.indirizzo_usato == s.indirizzo_completo):
            s.trasporto = t
            n += 1
        else:
            s.trasporto = None
    return n


COLONNA_MINUTI = "Minuti per tornare a casa"


def aggiorna_minuti_studenti(tabelle: dict[str, Tabella], trasporti: dict[str, Trasporto]) -> int:
    """Riempie le caselle vuote della colonna «Minuti per tornare a casa» (il migliore delle 4 fasce).

    Le caselle già scritte non si toccano: il numero messo a mano vince su quello calcolato, e per
    rifarlo calcolare basta svuotare la casella. Ritorna quante caselle ha riempito.
    """
    ts = tabelle.get(FOGLIO_STUDENTI)
    if ts is None or not ts.intestazione:
        return 0
    try:
        i_min = ts.colonna(COLONNA_MINUTI)
        i_cog, i_nome = ts.colonna("Cognome"), ts.colonna("Nome")
    except KeyError:
        return 0
    n = 0
    for r in ts.righe:
        while len(r) <= i_min:
            r.append("")
        cognome = (r[i_cog] or "").strip().upper()
        if not cognome:
            continue
        if (r[i_min] or "").strip():
            continue          # scritto a mano: non si tocca, ed è quello che vale per il calcolo
        t = trasporti.get(f"{cognome} {(r[i_nome] or '').strip().upper()}".strip())
        minuti = t.minuti_min if t is not None else None
        if minuti is not None:
            r[i_min] = str(minuti)
            n += 1
    return n


def aggiorna_totali_docenti(tabelle: dict[str, Tabella]) -> None:
    """Riempie la colonna «Ore dichiarate» e la riga TOTALE del foglio Docenti.

    Sono numeri di comodo per chi compila: il programma li ricalcola e non li legge mai.
    """
    from .template import RIGA_TOTALE

    td = tabelle.get(FOGLIO_DOCENTI)
    if td is None or not td.intestazione:
        return
    try:
        i_nome = td.colonna("Docente")
        i_tot = td.colonna("Ore dichiarate")
    except KeyError:
        return
    idx = []
    for nome_colonna in FASCE:
        try:
            idx.append(td.colonna(nome_colonna))
        except KeyError:
            return
    td.righe = [r for r in td.righe if r[i_nome].strip().upper() != RIGA_TOTALE]
    per_fascia = [0] * len(FASCE)
    totale = 0
    for r in td.righe:
        if not r[i_nome].strip():
            continue
        n = 0
        for k, c in enumerate(idx):
            if c < len(r) and (r[c] or "").strip():
                n += 1
                per_fascia[k] += 1
        r[i_tot] = str(n)
        totale += n
    if not td.righe:
        return
    riga = [""] * len(td.intestazione)
    riga[i_nome] = RIGA_TOTALE
    for k, c in enumerate(idx):
        riga[c] = str(per_fascia[k])
    riga[i_tot] = str(totale)
    td.righe.append(riga)


# ── Aggiornamento della struttura di file creati con versioni precedenti ─────

def _titolo_colonna(h: str) -> str:
    """Nome di colonna confrontabile: minuscolo, senza la parte fra parentesi."""
    return re.sub(r"\s*\(.*?\)", "", (h or "").strip().lower())


def aggiorna_struttura(tabelle: dict[str, Tabella]) -> list[str]:
    """Porta un file creato con una versione precedente alla struttura di oggi.

    Aggiunge i fogli e le colonne introdotti dopo, nella posizione giusta e senza toccare i dati,
    e le righe nuove del foglio Parametri. Ritorna l'elenco di quello che ha aggiunto.
    """
    from .template import (
        COLONNE_ABBINAMENTI, COLONNE_DOCENTI, COLONNE_GRUPPI, COLONNE_IMPEGNI, COLONNE_LMI,
        COLONNE_PARAMETRI, COLONNE_STUDENTI, PARAMETRI_DEFAULT,
    )

    modello = {FOGLIO_STUDENTI: COLONNE_STUDENTI, FOGLIO_IMPEGNI: COLONNE_IMPEGNI,
               FOGLIO_DOCENTI: COLONNE_DOCENTI, FOGLIO_GRUPPI: COLONNE_GRUPPI,
               FOGLIO_ABBINAMENTI: COLONNE_ABBINAMENTI, FOGLIO_LMI: COLONNE_LMI,
               FOGLIO_PARAMETRI: COLONNE_PARAMETRI}
    modifiche: list[str] = []

    # caso speciale: "Ore consecutive" è diventata "1° strumento attaccato" (SI = attaccate,
    # NO = in giorni diversi, vuoto = decide il programma). I SI scritti prima restano validi.
    ts = tabelle.get(FOGLIO_STUDENTI)
    nuova = COLONNE_STUDENTI[COLONNE_STUDENTI.index("Docente 2") + 1]
    if ts is not None and ts.intestazione:
        titoli = [_titolo_colonna(h) for h in ts.intestazione]
        if "ore consecutive" in titoli and _titolo_colonna(nuova) not in titoli:
            ts.intestazione[titoli.index("ore consecutive")] = nuova
            modifiche.append(f"colonna «Ore consecutive» rinominata «{nuova}» nel foglio {FOGLIO_STUDENTI}")

    # caso speciale: l'aula unica per tutta la settimana è diventata una per giorno
    td = tabelle.get(FOGLIO_DOCENTI)
    if td is not None and td.intestazione:
        presenti = [h.strip().lower() for h in td.intestazione]
        if "aula" in presenti and "aula lun" not in presenti:
            i_aula = presenti.index("aula")
            td.intestazione[i_aula] = COLONNE_AULE[0]
            for k, nome_col in enumerate(COLONNE_AULE[1:], start=1):
                td.intestazione.insert(i_aula + k, nome_col)
                for r in td.righe:
                    r.insert(i_aula + k, r[i_aula] if i_aula < len(r) else "")
            modifiche.append(f"colonne {', '.join(COLONNE_AULE)} nel foglio {FOGLIO_DOCENTI} "
                             "(prima l'aula era una sola per tutta la settimana)")

    for nome_foglio, colonne_modello in modello.items():
        tab = tabelle.get(nome_foglio)
        if tab is None:
            continue
        if not tab.intestazione:            # foglio nuovo, che la lettura ha creato vuoto
            tab.intestazione = list(colonne_modello)
            modifiche.append(f"foglio «{nome_foglio}»")
            continue
        presenti = [_titolo_colonna(h) for h in tab.intestazione]
        for colonna in colonne_modello:
            atteso = _titolo_colonna(colonna)
            if atteso in presenti:
                continue
            pos = len(tab.intestazione)     # in coda, se non si trova un aggancio
            for prec in reversed(colonne_modello[:colonne_modello.index(colonna)]):
                if _titolo_colonna(prec) in presenti:
                    pos = presenti.index(_titolo_colonna(prec)) + 1
                    break
            tab.intestazione.insert(pos, colonna)
            presenti.insert(pos, atteso)
            for r in tab.righe:
                r.insert(pos, "")
            modifiche.append(f"colonna «{colonna}» nel foglio {nome_foglio}")

    tp = tabelle.get(FOGLIO_PARAMETRI)
    if tp is not None and tp.intestazione:
        chiavi = [tp.valore(r, "Parametro").lower() for r in tp.righe]
        for nome, valore, spiegazione in PARAMETRI_DEFAULT:
            if nome.lower() not in chiavi:
                tp.righe.append([nome, str(valore), spiegazione][:len(tp.intestazione)])
                modifiche.append(f"riga «{nome}» nel foglio {FOGLIO_PARAMETRI}")
    return modifiche

