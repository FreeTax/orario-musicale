"""Crea un file di input di esempio, con dati INVENTATI, pronto da calcolare.

Serve per provare il programma senza usare i dati veri della scuola: 30 ragazzi di fantasia
in cinque classi, 7 docenti con disponibilità già compilate, 6 gruppi di musica da camera,
2 laboratori LMI, qualche eccezione (giorno unico, ore consecutive, giorni non disponibili)
e un docente con ore di pianista accompagnatore.

Gli indirizzi sono vie realmente esistenti nei comuni intorno a Pistoia, così funziona anche
"Aggiorna trasporti"; i nomi delle persone sono inventati.

Uso:
    .venv/bin/python dati_esempio.py [percorso_file.xlsx]
"""

from __future__ import annotations

import sys
from collections import Counter
from pathlib import Path

from openpyxl.styles import PatternFill

from orario.costanti import FASCE, FOGLIO_DOCENTI, FOGLIO_LMI, FOGLIO_PARAMETRI, FOGLIO_STUDENTI, N_ORE
from orario.template import COLONNE_GRUPPI, ESEMPIO, costruisci_workbook

USCITA_DEFAULT = Path(__file__).parent / "esempio" / "orario_2026-27.xlsx"
INDIRIZZO_SCUOLA = "Corso Antonio Gramsci 148, Pistoia"

# ── Docenti: nome, strumenti, aula, ore di accompagnamento ───────────────────
DOCENTI = [
    ("Bianchi", ["PIANOFORTE"], "M1", 2),
    ("Conti", ["PIANOFORTE"], "M2", 0),
    ("Verdi", ["VIOLINO", "VIOLONCELLO"], "M3", 0),
    ("Ferrari", ["CHITARRA"], "M4", 0),
    ("Marini", ["CANTO"], "M5", 0),
    ("Galli", ["FLAUTO", "CLARINETTO", "SAXOFONO"], "M6", 0),
    ("Lombardi", ["PERCUSSIONI"], "M7", 0),
]

# ── Studenti: classe, cognome, nome, comune, via, civico, km, 1° strum., 2° strum. ──
# (la 5ª non ha il 2° strumento; il docente lo assegna lo script bilanciando le ore)
STUDENTI = [
    # classe 1ª — 8 ragazzi
    (1, "AMATO", "GIULIA", "Pistoia", "Via Antonio Pacinotti", "12", 1.2, "PIANOFORTE", "VIOLINO"),
    (1, "BRUNI", "MATTEO", "Quarrata", "Via Firenze", "45", 12.0, "CHITARRA", "PIANOFORTE"),
    (1, "CARRARA", "SOFIA", "Agliana", "Via Provinciale Pratese", "88", 11.5, "CANTO", "PIANOFORTE"),
    (1, "DONATI", "LORENZO", "Pistoia", "Via Dalmazia", "220", 2.5, "FLAUTO", "PIANOFORTE"),
    (1, "ESPOSITO", "CHIARA", "Serravalle Pistoiese", "Via Provinciale Lucchese", "30", 6.5, "PIANOFORTE", "CANTO"),
    (1, "FANTI", "TOMMASO", "Montale", "Via Martiri della Libertà", "17", 8.0, "PERCUSSIONI", "PIANOFORTE"),
    (1, "GRASSI", "ALICE", "Pistoia", "Via Bonellina", "5", 3.0, "VIOLINO", "PIANOFORTE"),
    (1, "LUPI", "FEDERICO", "Monsummano Terme", "Via Francesca", "102", 14.0, "CLARINETTO", "PIANOFORTE"),
    # classe 2ª — 6 ragazzi
    (2, "MARCHI", "EMMA", "Pistoia", "Via Ciliegiole", "40", 4.0, "PIANOFORTE", "FLAUTO"),
    (2, "NARDI", "GABRIELE", "Larciano", "Via Corsini", "210", 23.0, "SAXOFONO", "PIANOFORTE"),
    (2, "ORLANDI", "VIOLA", "Pistoia", "Via Sestini", "8", 1.8, "CANTO", "PIANOFORTE"),
    (2, "PIERI", "LEONARDO", "Quarrata", "Via Vecchia Fiorentina", "60", 13.5, "VIOLONCELLO", "PIANOFORTE"),
    (2, "RENZI", "MARTINA", "Pieve a Nievole", "Via Francesca Vecchia", "22", 15.0, "PIANOFORTE", "VIOLINO"),
    (2, "SARTI", "DIEGO", "Pistoia", "Via del Bonellina", "77", 3.5, "PERCUSSIONI", "CHITARRA"),
    # classe 3ª — 6 ragazzi (1° + 2° + musica da camera)
    (3, "TOSI", "BEATRICE", "Pistoia", "Via Antonio Gramsci", "20", 0.9, "PIANOFORTE", "CANTO"),
    (3, "VESTRI", "SAMUELE", "Sambuca Pistoiese", "Località Pavana", "14", 28.0, "CHITARRA", "PIANOFORTE"),
    (3, "ZACCARIA", "NOEMI", "Pescia", "Via Amendola", "35", 30.0, "FLAUTO", "PIANOFORTE"),
    (3, "BANDINI", "FILIPPO", "Pistoia", "Via Puccini", "3", 1.5, "VIOLINO", "PIANOFORTE"),
    (3, "CECCONI", "GAIA", "Montecatini Terme", "Viale Verdi", "50", 16.0, "CANTO", "PIANOFORTE"),
    (3, "DINI", "RICCARDO", "Pistoia", "Via del Villone", "11", 2.8, "PIANOFORTE", "PERCUSSIONI"),
    # classe 4ª — 5 ragazzi
    (4, "FIORI", "ELISA", "San Marcello Piteglio", "Via Giacomo Matteotti", "9", 24.0, "PIANOFORTE", "FLAUTO"),
    (4, "GARGINI", "NICOLA", "Pistoia", "Via Fiorentina", "150", 3.2, "VIOLONCELLO", "PIANOFORTE"),
    (4, "LAZZERI", "ARIANNA", "Lamporecchio", "Via Roma", "18", 26.0, "CANTO", "CHITARRA"),
    (4, "MERLI", "ANDREA", "Pistoia", "Via Bure Vecchia Nord", "64", 5.0, "CHITARRA", "PIANOFORTE"),
    (4, "NESTI", "GIORGIA", "Quarrata", "Via Montalbano", "12", 14.5, "CLARINETTO", "PIANOFORTE"),
    # classe 5ª — 5 ragazzi (2 ore di 1° strumento + musica da camera)
    (5, "PAGNI", "DAVIDE", "Pistoia", "Via Vecchia Pratese", "31", 4.5, "PIANOFORTE", ""),
    (5, "QUERCI", "SARA", "Cantagallo", "Via Cavallaie", "7", 31.0, "VIOLINO", ""),
    (5, "SALVI", "MARCO", "Pistoia", "Via San Biagio", "26", 2.0, "CHITARRA", ""),
    (5, "TURCHI", "LUCIA", "Buggiano", "Via Roma", "10", 20.0, "CANTO", ""),
    (5, "VENTURI", "PIETRO", "Pistoia", "Via Copernico", "15", 6.0, "PERCUSSIONI", ""),
]

# ── Gruppi di musica da camera: docente e cognomi (classi 3ª, 4ª, 5ª) ────────
GRUPPI = [
    ("Verdi", ["Bandini", "Gargini", "Querci"]),
    ("Ferrari", ["Vestri", "Merli", "Salvi"]),
    ("Marini", ["Cecconi", "Lazzeri", "Turchi"]),
    ("Galli", ["Zaccaria", "Nesti", "Fiori"]),
    ("Lombardi", ["Dini", "Venturi"]),
    ("Conti", ["Tosi", "Pagni"]),
]

# ── Eccezioni per singolo ragazzo ────────────────────────────────────────────
ORE_CONSECUTIVE = {"AMATO", "QUERCI", "MARCHI"}       # le 2 ore di 1° strumento attaccate
GIORNO_UNICO = {"VESTRI", "ZACCARIA"}                  # un solo rientro a settimana
GIORNI_NON_DISPONIBILI = {"LUPI": "Mar", "NESTI": "Gio", "TURCHI": "Lun, Ven"}

# ── Laboratori LMI: solo ricopiati nell'orario, non calcolati ────────────────
LMI = [
    ["LMI ARCHI E PIANOFORTE", "1ª e 2ª", "Verdi", "Aula M3, 1° piano", "Martedì 4ª-5ª ora",
     "Amato, Grassi, Renzi, Pieri", ""],
    ["LMI FIATI E PERCUSSIONI", "1ª e 2ª", "Galli", "Aula M6, 1° piano", "Giovedì 4ª-5ª ora",
     "Donati, Lupi, Fanti, Sarti, Nardi", ""],
]

# quante fasce lasciare libere a ogni docente, oltre alle ore che gli servono
MARGINE_FASCE = 4


def main(uscita: Path) -> None:
    # ── docente per ogni lezione, bilanciando le ore ──
    per_strumento: dict[str, list[str]] = {}
    for nome, strumenti, _, _ in DOCENTI:
        for s in strumenti:
            per_strumento.setdefault(s, []).append(nome)
    ore: Counter = Counter()
    for nome, _, _, accomp in DOCENTI:
        ore[nome] += accomp
    for docente, _ in GRUPPI:
        ore[docente] += 1

    mancanti: set[str] = set()

    def scegli(strumento: str, quante: int) -> str:
        candidati = per_strumento.get(strumento, [])
        if not candidati:
            mancanti.add(strumento)
            return ""
        d = min(candidati, key=lambda x: ore[x])
        ore[d] += quante
        return d

    righe_studenti = []
    for classe, cognome, nome, comune, via, civico, km, s1, s2 in STUDENTI:
        h1 = 2 if classe in (1, 2, 5) else 1
        h2 = 0 if classe == 5 else 1
        righe_studenti.append({
            "classe": classe, "cognome": cognome, "nome": nome, "comune": comune,
            "indirizzo": via, "civico": civico, "km": km,
            "strum1": s1, "doc1": scegli(s1, h1),
            "strum2": s2, "doc2": scegli(s2, h2) if s2 else "",
        })
    if mancanti:
        raise SystemExit(f"Nessun docente per: {', '.join(sorted(mancanti))}. Aggiungerlo in DOCENTI.")

    # ── disponibilità: giorni pieni, quanti bastano per le ore + margine ──
    righe_docenti = []
    for nome, strumenti, aula, accomp in DOCENTI:
        servono = ore[nome]
        giorni = min(5, max(2, -(-(servono + MARGINE_FASCE) // N_ORE)))
        # aula fissa tutta la settimana, tranne un docente che il mercoledì e il giovedì cambia stanza
        aule = [aula] * 5
        if nome == "Verdi":
            aule[2] = aule[3] = "M8"
        righe_docenti.append({
            "nome": nome, "strumenti": ", ".join(strumenti), "aule": aule, "ore_accomp": accomp,
            "giorni": giorni, "ore": servono,
        })

    wb = costruisci_workbook(righe_studenti, righe_docenti, esempi=False)

    # ── eccezioni e disponibilità nelle celle ──
    ws = wb[FOGLIO_STUDENTI]
    intest = [c.value for c in ws[1]]
    i_cog, i_cons = intest.index("Cognome"), intest.index("Ore consecutive (SI/NO)")
    i_unico, i_giorni = intest.index("Giorno unico (SI/NO)"), intest.index("Giorni NON disponibili")
    for r in range(2, ws.max_row + 1):
        cog = ws.cell(row=r, column=i_cog + 1).value
        if cog in ORE_CONSECUTIVE:
            ws.cell(row=r, column=i_cons + 1).value = "SI"
        if cog in GIORNO_UNICO:
            ws.cell(row=r, column=i_unico + 1).value = "SI"
        if cog in GIORNI_NON_DISPONIBILI:
            ws.cell(row=r, column=i_giorni + 1).value = GIORNI_NON_DISPONIBILI[cog]

    ws = wb[FOGLIO_DOCENTI]
    intest = [c.value for c in ws[1]]
    col_prima = intest.index(FASCE[0]) + 1
    for r, d in enumerate(righe_docenti, start=2):
        for g in range(d["giorni"]):
            for o in range(N_ORE):
                ws.cell(row=r, column=col_prima + g * N_ORE + o).value = "X"

    # ── gruppi di musica da camera ──
    ws = wb["Gruppi LMC"]
    for n, (docente, membri) in enumerate(GRUPPI, start=1):
        ws.append([n, docente] + membri + [""] * (5 - len(membri)) + [""])

    # ── laboratori LMI ──
    ws = wb[FOGLIO_LMI]
    for riga in LMI:
        ws.append(riga)

    # ── indirizzo della scuola ──
    ws = wb[FOGLIO_PARAMETRI]
    for r in range(2, ws.max_row + 1):
        if "scuola" in str(ws.cell(row=r, column=1).value or "").lower():
            ws.cell(row=r, column=2).value = INDIRIZZO_SCUOLA

    uscita.parent.mkdir(parents=True, exist_ok=True)
    wb.save(uscita)

    print(f"Creato: {uscita}")
    print(f"  {len(STUDENTI)} studenti (per classe: "
          + ", ".join(f"{c}ª={sum(1 for s in STUDENTI if s[0] == c)}" for c in range(1, 6)) + ")")
    print(f"  {len(DOCENTI)} docenti, {len(GRUPPI)} gruppi di musica da camera, {len(LMI)} laboratori LMI")
    print("  ore e disponibilità per docente:")
    for d in righe_docenti:
        print(f"    {d['nome']:10s} {d['ore']:2d} ore su {d['giorni'] * N_ORE:2d} fasce disponibili"
              + (f"  (di cui {d['ore_accomp']} di accompagnamento)" if d["ore_accomp"] else ""))
    print(f"  eccezioni: ore consecutive {sorted(ORE_CONSECUTIVE)}, giorno unico {sorted(GIORNO_UNICO)}, "
          f"giorni vietati {GIORNI_NON_DISPONIBILI}")


if __name__ == "__main__":
    main(Path(sys.argv[1]) if len(sys.argv) > 1 else USCITA_DEFAULT)
